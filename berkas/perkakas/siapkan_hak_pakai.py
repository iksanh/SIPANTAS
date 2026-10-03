# -*- coding: utf-8 -*-
"""Bangun tata naskah Hak Pakai (bap-hp, risalah-hp, sk-hp) dari template Hak Milik.

Tata naskah Hak Pakai Kantah Bone Bolango (contohnya di folder `hak_pakai/`)
sama kerangkanya dengan Hak Milik: kop, susunan panitia, Data Pendukung, dasar
hukum, dan diktum SK-nya sama. Yang berbeda:

- sebutan haknya «Hak Pakai dengan jangka waktu» dan jangka waktunya disebut
  (Pasal 139 ayat (4) Permen ATR/BPN 18/2021);
- subjeknya boleh perorangan maupun badan hukum (Pasal 111 ayat (2)). Bagian
  yang berbeda antara keduanya mengikuti format Lampiran VI (SK) dan Lampiran
  VII (Risalah) Permen 18/2021, dijadikan bersyarat dalam satu template:
  `{{!badan_hukum}}` untuk perorangan, `{{?badan_hukum}}` untuk badan hukum;
- diktum KEDUA menambah butir perpanjangan Hak Pakai.

Karena itu template Hak Pakai **dirakit dari template Hak Milik yang sedang
terpasang**, bukan dari dokumen Word tersendiri — kop dan gaya hurufnya ikut
apa adanya. Setiap perubahan mencari kalimat jangkarnya di template Hak Milik;
kalau jangkarnya tidak ketemu (template Hak Milik sudah disunting jauh),
perkakas ini berhenti dan menyebut kalimat mana, bukan menghasilkan template
setengah jadi.

Sesudah terbit, ketiganya jadi template biasa di menu Template → tab Hak Pakai:
diunduh, disunting di Word, dan diunggah lagi. Menjalankan perkakas ini lagi
**menimpa** suntingan itu:

    python -m berkas.perkakas.siapkan_hak_pakai
"""
import copy
import os
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .. import jalur


class JangkarHilang(Exception):
    """Kalimat yang hendak diubah tidak ditemukan di template Hak Milik."""


# ------------------------------------------------------------------ paragraf
def _teks(p):
    return "".join(r.text for r in p.runs)


def _semua_paragraf(wadah):
    for p in wadah.paragraphs:
        yield p
    for t in wadah.tables:
        for baris in t.rows:
            for sel in baris.cells:
                yield from _semua_paragraf(sel)


def _ganti_run(p, lama, baru):
    """Ganti `lama` dengan `baru` di dalam run, walau `lama` terpecah ke beberapa
    run. Format run tempat awal kecocokan yang dipakai. Kembalikan jumlahnya."""
    n = 0
    while True:
        runs = list(p.runs)
        teks = "".join(r.text for r in runs)
        i = teks.find(lama)
        if i < 0 or not lama:
            return n
        peta = [(ri, j) for ri, r in enumerate(runs) for j in range(len(r.text))]
        r_awal, o_awal = peta[i]
        r_akhir, o_akhir = peta[i + len(lama) - 1]
        if r_awal == r_akhir:
            s = runs[r_awal].text
            runs[r_awal].text = s[:o_awal] + baru + s[o_akhir + 1:]
        else:
            runs[r_awal].text = runs[r_awal].text[:o_awal] + baru
            for ri in range(r_awal + 1, r_akhir):
                runs[ri].text = ""
            runs[r_akhir].text = runs[r_akhir].text[o_akhir + 1:]
        n += 1
        if lama in baru:                      # jangan berputar tanpa ujung
            return n


def _ganti_semua(dok, lama, baru):
    return sum(_ganti_run(p, lama, baru) for p in _semua_paragraf(dok))


def _cari(dok, jangkar, ke=0):
    """Paragraf badan dokumen yang memuat `jangkar` (yang ke-`ke` bila banyak)."""
    ketemu = [p for p in dok.paragraphs if jangkar in _teks(p)]
    if len(ketemu) <= ke:
        raise JangkarHilang(jangkar)
    return ketemu[ke]


def _ganti_di(dok, jangkar, lama, baru):
    p = _cari(dok, jangkar)
    if not _ganti_run(p, lama, baru):
        raise JangkarHilang(f"{lama!r} di paragraf «{jangkar}»")
    return p


def _setel(p, teks):
    """Isi paragraf dengan `teks` memakai format run pertamanya."""
    runs = [r for r in p.runs]
    if not runs:
        p.add_run(teks)
        return
    utama = next((r for r in runs if r.text.strip()), runs[0])
    for r in runs:
        r.text = ""
    utama.text = teks


def _awali(p, penanda):
    """Taruh penanda syarat di depan teks paragraf (di run pertama yang berisi)."""
    for r in p.runs:
        if r.text:
            r.text = penanda + r.text
            return
    p.add_run(penanda)


def _salin_sesudah(p, teks):
    """Paragraf baru sesudah `p`, berformat sama (termasuk penomorannya)."""
    salinan = copy.deepcopy(p._p)
    p._p.addnext(salinan)
    baru = Paragraph(salinan, p._parent)
    _setel(baru, teks)
    return baru


def _salin_sebelum(p, teks):
    salinan = copy.deepcopy(p._p)
    p._p.addprevious(salinan)
    baru = Paragraph(salinan, p._parent)
    _setel(baru, teks)
    return baru


# --------------------------------------------------------------------- tabel
def _tabel_dengan(dok, jangkar):
    for el in dok.element.body.iterchildren():
        if el.tag == qn("w:tbl"):
            t = Table(el, dok)
            if jangkar in "".join(_teks(p) for p in _semua_paragraf_tabel(t)):
                return t
    raise JangkarHilang(f"tabel yang memuat «{jangkar}»")


def _semua_paragraf_tabel(t):
    for baris in t.rows:
        for sel in baris.cells:
            yield from sel.paragraphs


def _sel(tr):
    return tr.findall(qn("w:tc"))


def _setel_sel(tc, teks):
    ps = tc.findall(qn("w:p"))
    for i, p_el in enumerate(ps):
        p = Paragraph(p_el, None)
        if i == 0:
            _setel(p, teks)
        else:
            _setel(p, "")


def _tiru_baris(tr_contoh, isi):
    """Baris baru tiruan `tr_contoh`, belum ditempatkan. `isi` = teks tiap sel;
    None berarti selnya dibiarkan seperti contohnya."""
    tr = copy.deepcopy(tr_contoh)
    for tc, teks in zip(_sel(tr), isi):
        if teks is not None:
            _setel_sel(tc, teks)
    return tr


def _baris_penanda(tr_contoh, penanda):
    """Baris yang hanya memuat penanda blok; dibuang docxgen saat merakit."""
    return _tiru_baris(tr_contoh, [penanda] + [""] * (len(_sel(tr_contoh)) - 1))


def _tiru_paragraf(contoh, sesudah, teks):
    """Paragraf tiruan `contoh` (format dan penomorannya) sesudah `sesudah`."""
    el = copy.deepcopy(contoh._p)
    sesudah._p.addnext(el)
    baru = Paragraph(el, contoh._parent)
    _setel(baru, teks)
    return baru


# ---------------------------------------------------------------- kalimat baku
# Telaah subjek. Rujukan pasalnya: UUPA Pasal 42, PP 18/2021 Pasal 49 ayat (2),
# Permen ATR/BPN 18/2021 Pasal 111 ayat (2) — huruf a WNI, huruf b badan hukum.
TELAAH_BH = (
    "Berdasarkan Undang-Undang Nomor 5 Tahun 1960 tentang Peraturan Dasar Pokok-Pokok "
    "Agraria Pasal 42 huruf c, Peraturan Pemerintah Nomor 18 Tahun 2021 tentang Hak "
    "Pengelolaan, Hak atas Tanah, Satuan Rumah Susun, dan Pendaftaran Tanah Pasal 49 ayat (2) "
    "huruf b, dan Peraturan Menteri Agraria dan Tata Ruang/Kepala Badan Pertanahan Nasional "
    "Nomor 18 Tahun 2021 Tentang Tata Cara Penetapan Hak Pengelolaan Dan Hak Atas Tanah "
    "Pasal 111 ayat (2) huruf b, bahwa pemohon sebagai badan hukum yang didirikan menurut "
    "hukum Indonesia dan berkedudukan di Indonesia memenuhi syarat sebagai subyek "
    "{{nama_hak_lengkap}}.")

ANALISIS_BH = (
    "Bahwa {{bh_uraian_subjek}}{{bh_diwakili}}, sehingga berdasarkan Undang-Undang Nomor 5 "
    "Tahun 1960 tentang Peraturan Dasar Pokok-Pokok Agraria Pasal 42 huruf c dan Peraturan "
    "Menteri Agraria dan Tata Ruang/Kepala Badan Pertanahan Nasional Nomor 18 Tahun 2021 "
    "Tentang Tata Cara Penetapan Hak Pengelolaan Dan Hak Atas Tanah Pasal 111 ayat (2) "
    "huruf b, pemohon memenuhi syarat sebagai subyek {{nama_hak_lengkap}}.")

# Rujukan subjek perorangan pada template Hak Milik -> rujukan Hak Pakai
PASAL_PERORANGAN = [
    ("pasal 42 ayat (1)", "Pasal 42 huruf a"),
    ("Pasal 52 ayat (1) huruf a", "Pasal 111 ayat (2) huruf a"),
]

DATA_PENDUKUNG_BH = [
    "{{?badan_hukum}}{{?bh_akta_uraian}}Fotokopi {{bh_akta_uraian}} tentang pendirian "
    "{{nama_penerima}};",
    "{{?badan_hukum}}{{?bh_perubahan_uraian}}Fotokopi {{bh_perubahan_uraian}} tentang "
    "perubahan anggaran dasar {{nama_penerima}};",
    "{{?badan_hukum}}{{?bh_pengesahan_ringkas}}Fotokopi {{bh_pengesahan_ringkas}} tentang "
    "pengesahan badan hukum {{nama_penerima}};",
    "{{?badan_hukum}}{{?bh_nib}}Fotokopi Nomor Induk Berusaha {{bh_nib}} atas nama "
    "{{nama_penerima}};",
    "{{?badan_hukum}}{{?bh_wakil_nama}}Fotokopi KTP {{bh_wakil_nama}} selaku "
    "{{bh_wakil_jabatan}} {{nama_penerima}};",
    "{{?badan_hukum}}{{?bh_csr}}Salinan {{bh_csr}} tentang kesanggupan melaksanakan "
    "tanggung jawab sosial dan lingkungan;",
]

JANGKA = "{{jangka_tahun}} ({{jangka_terbilang}}) tahun"


# ------------------------------------------------------------------------ BAP
def ubah_bap(dok):
    # «Sdr. PT Maju Jaya» tidak lazim: badan hukum disebut namanya saja
    _ganti_di(dok, "Dengan ini kami telah melakukan pemeriksaan lapang",
              "{{sapaan}} {{nama_penerima}}", "{{penerima_disapa}}")


# -------------------------------------------------------------------- Risalah
def ubah_risalah(dok):
    _ganti_semua(dok, "{{nama_hak}}", "{{nama_hak_lengkap}}")
    _ganti_semua(dok, "Hak Milik", "{{nama_hak_lengkap}}")

    # I. Uraian mengenai pemohon: baris perorangan atau badan hukum
    # Susunannya jadi: {{#pemohon_perorangan}} 1. Perorangan a-e {{/…}}
    #                  {{#pemohon_badan_hukum}} 1. Badan Hukum a-e {{/…}}
    t = _tabel_dengan(dok, "Nama Pemohon")
    baris = t._tbl.findall(qn("w:tr"))
    kepala, contoh = baris[0], baris[1]
    if "Perorangan" not in "".join(x.text or "" for x in kepala.iter(qn("w:t"))):
        raise JangkarHilang("baris «1. Perorangan» pada tabel Uraian mengenai Pemohon")
    kepala.addprevious(_baris_penanda(contoh, "{{#pemohon_perorangan}}"))
    urutan = [_baris_penanda(contoh, "{{/pemohon_perorangan}}"),
              _baris_penanda(contoh, "{{#pemohon_badan_hukum}}")]
    kepala_bh = copy.deepcopy(kepala)
    for tc in _sel(kepala_bh):
        if "Perorangan" in "".join(x.text or "" for x in tc.iter(qn("w:t"))):
            _setel_sel(tc, "Badan Hukum")
    urutan.append(kepala_bh)
    n_sel = len(_sel(contoh))
    for huruf, label, nilai in (
            ("a.", "Nama Pemohon", "{{nama_penerima}}"),
            ("b.", "Domisili/Tempat Kedudukan", "{{bh_kedudukan}}"),
            ("c.", "Akta Pendirian Badan Hukum/Peraturan Pendirian Perusahaan",
             "{{bh_akta_uraian}}"),
            ("d.", "Pengesahan Badan Hukum", "{{bh_pengesahan_ringkas}}"),
            ("e.", "Nomor Induk Berusaha/Tanda Daftar Perusahaan", "{{bh_nib}}")):
        urutan.append(_tiru_baris(contoh, [None] * (n_sel - 4) + [huruf, label, ":", nilai]))
    urutan.append(_baris_penanda(contoh, "{{/pemohon_badan_hukum}}"))
    sebelum = baris[-1]
    for tr in urutan:
        sebelum.addnext(tr)
        sebelum = tr

    # III. Uraian atas hak: jangka waktu
    t = _tabel_dengan(dok, "Jangka waktu")
    for b in t.rows:
        if "Jangka waktu" in "".join(c.text for c in b.cells):
            _setel_sel(b._tr.findall(qn("w:tc"))[-1], "{{jangka_tahun}} Tahun")
            break

    # IV. Data pendukung: KTP pemohon hanya untuk perorangan, lalu surat badan hukum
    ktp = _cari(dok, "Fotokopi KTP dan Kartu Keluarga {{nama_penerima}}")
    _awali(ktp, "{{!badan_hukum}}")
    sebelum = ktp
    for kalimat in DATA_PENDUKUNG_BH:
        sebelum = _salin_sesudah(sebelum, kalimat)
    _awali(_cari(dok, "Fotokopi KTP dan Kartu Keluarga {{pemilik_sebelumnya}}"),
           "{{?pemilik_sebelumnya}}")

    # VI. Telaah subjek: perorangan atau badan hukum
    judul_po = _cari(dok, "URAIAN DAN TELAAH ATAS SUBYEK HAK")
    po = [Paragraph(el, judul_po._parent) for el in
          (judul_po._p.getnext(), judul_po._p.getnext().getnext(),
           judul_po._p.getnext().getnext().getnext())]
    if _teks(po[0]).strip() != "Perorangan":
        raise JangkarHilang("paragraf «Perorangan» sesudah judul VI")
    for lama, baru in PASAL_PERORANGAN:
        _ganti_run(po[2], lama, baru)
    for p in po:
        _awali(p, "{{!badan_hukum}}")
    # tiruan ketiganya untuk badan hukum: judul, uraian, telaah - penomorannya ikut
    sebelum = po[2]
    for contoh, teks in ((po[0], "Badan Hukum"),
                         (po[1], "{{bh_uraian_subjek}}{{bh_diwakili}}."),
                         (po[2], TELAAH_BH)):
        sebelum = _tiru_paragraf(contoh, sebelum, "{{?badan_hukum}}" + teks)

    # VIII. Analisis: kalimat subjek
    analisis = _cari(dok, "tempat tanggal lahir {{ttl_penerima}} Kewarganegaraan Indonesia, "
                          "bertempat tinggal di {{alamat_penerima}}, Kab. Bone Bolango")
    _ganti_run(analisis, ", Kab. Bone Bolango pekerjaan", " pekerjaan")
    for lama, baru in PASAL_PERORANGAN:
        _ganti_run(analisis, lama, baru)
    _awali(analisis, "{{!badan_hukum}}")
    _salin_sesudah(analisis, "{{?badan_hukum}}" + ANALISIS_BH)

    # X. Kesimpulan: jangka waktu yang direkomendasikan (Pasal 139 ayat (4))
    _ganti_di(dok, "Berdasarkan uraian tersebut di atas (poin 1 sampai dengan poin 4)",
              "untuk diberikan {{nama_hak_lengkap}}.",
              f"untuk diberikan {{{{nama_hak_lengkap}}}} selama {JANGKA}.")


# ------------------------------------------------------------------------- SK
def ubah_sk(dok):
    _ganti_semua(dok, "{{nama_hak_kapital}}", "{{nama_hak_lengkap_kapital}}")
    _ganti_semua(dok, "{{nama_hak}}", "{{nama_hak_lengkap_judul}}")
    _ganti_semua(dok, "Hak Milik", "{{nama_hak_lengkap_judul}}")

    # Membaca a: badan hukum «berkedudukan di», bukan «bertempat tinggal di»
    _ganti_di(dok, "surat permohonan {{nama_hak_lengkap_judul}} tanggal",
              "bertempat tinggal di {{alamat_penerima}}", "{{domisili_penerima}}")

    # Menimbang a: subjek perorangan / badan hukum (Lampiran VI format A.1.a/A.1.b)
    a = _cari(dok, "adalah Warga Negara Indonesia, bertempat tinggal di {{alamat_penerima}}")
    # alamat pemohon sudah lengkap sampai provinsi; kabupaten tidak ditempel lagi
    _ganti_run(a, "{{alamat_penerima}},Kabupaten Bone Bolango pemegang",
               "{{alamat_penerima}}, pemegang")
    _awali(a, "{{!badan_hukum}}")
    _salin_sesudah(a, "{{?badan_hukum}}bahwa {{bh_uraian_subjek}}, sehingga telah memenuhi "
                      "syarat sebagai subjek {{nama_hak_lengkap_judul}};")

    # Menimbang d.6: berkesimpulan ... dengan jangka waktu
    _ganti_di(dok, "berkesimpulan permohonan tersebut dapat dipertimbangkan",
              "diberikan {{nama_hak_lengkap_judul}} kepada {{nama_penerima}} ,",
              f"diberikan {{{{nama_hak_lengkap_judul}}}} selama {JANGKA} kepada "
              "{{nama_penerima}},")

    # Menimbang: CSR badan hukum usaha sumber daya alam, sebelum huruf BPHTB
    _salin_sebelum(_cari(dok, "bahwa berdasarkan Undang-Undang Nomor 1 Tahun 2022"),
                   "{{?bh_csr}}bahwa dalam rangka melaksanakan tanggung jawab sosial dan "
                   "lingkungan (Corporate Social Responsibility), pemohon telah menyatakan "
                   "kesanggupannya untuk melaksanakan tanggung jawab sosial dan lingkungan "
                   "sebagaimana {{bh_csr}};")

    # Mengingat: undang-undang bentuk badan hukumnya, sesudah UU Penataan Ruang
    _salin_sesudah(_cari(dok, "Undang-Undang Nomor 26 Tahun 2007 tentang Penataan Ruang"),
                   "{{?bh_dasar_bentuk}}{{bh_dasar_bentuk}};")

    # Menetapkan dan KESATU menyebut jangka waktunya
    _ganti_di(dok, "Menetapkan", "PEMBERIAN {{nama_hak_lengkap_kapital}} ATAS NAMA",
              "PEMBERIAN {{nama_hak_lengkap_kapital}} SELAMA {{jangka_tahun}} "
              "({{jangka_terbilang_kapital}}) TAHUN ATAS NAMA")
    _ganti_di(dok, "Memberikan kepada {{nama_penerima}}",
              "bertempat tinggal di {{alamat_penerima}}, {{nama_hak_lengkap_kapital}} atas",
              f"{{{{domisili_penerima}}}}, {{{{nama_hak_lengkap_kapital}}}} selama {JANGKA} atas")

    # KEDUA: butir perpanjangan Hak Pakai dengan jangka waktu (Pasal 113 ayat (1))
    _salin_sesudah(_cari(dok, "penerima hak diwajibkan untuk menyerahkan kembali tanah"),
                   "{{?jangka_tahun}}Hak Pakai ini dapat diperpanjang dengan jangka waktu "
                   "paling lama 20 (dua puluh) tahun apabila menurut penilaian Pemerintah "
                   "tanah Hak Pakai ini telah diusahakan dengan baik dan memenuhi "
                   "syarat-syarat yang ditentukan peraturan perundang-undangan;")


UBAH = {"bap": ubah_bap, "risalah": ubah_risalah, "sk": ubah_sk}


def siapkan(jenis, asal, tujuan):
    dok = Document(asal)
    UBAH[jenis](dok)
    dok.save(tujuan)
    return tujuan


def main():
    gagal = False
    for jenis in UBAH:
        asal = os.path.join(jalur.TEMPLATE, f"{jenis}.docx")
        tujuan = os.path.join(jalur.TEMPLATE, f"{jenis}-hp.docx")
        try:
            siapkan(jenis, asal, tujuan)
        except JangkarHilang as ex:
            gagal = True
            print(f"  ! {jenis}: kalimat jangkar tidak ditemukan di {jenis}.docx: {ex}")
            continue
        print(f"  dibuat  {tujuan}")
    return 1 if gagal else 0


if __name__ == "__main__":
    sys.exit(main())
