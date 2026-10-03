# -*- coding: utf-8 -*-
"""Bangun template DOCX untuk daftar kelengkapan dan surat pengembalian berkas.

Berbeda dari `siapkan_template.py`, yang mengubah dokumen Word ber-MERGEFIELD
jadi template sistem. Dua formulir di sini tidak punya dokumen sumber: bentuknya
ditetapkan Lampiran Permen ATR/BPN 18/2021, jadi kerangkanya dibangun sekali di
sini lalu diperlakukan seperti template lain — diunduh, disunting di Word, dan
diunggah lagi lewat menu Template.

Aman diulang, tetapi **menimpa** hasil suntingan. Jalankan hanya kalau
templatenya memang mau dikembalikan ke bentuk awal:

    python -m berkas.perkakas.siapkan_checklist
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .. import jalur

HURUF = "Times New Roman"
UKURAN = Pt(11)


def _atur_gaya(dok):
    g = dok.styles["Normal"]
    g.font.name = HURUF
    g.font.size = UKURAN
    g.paragraph_format.space_after = Pt(0)
    g.paragraph_format.space_before = Pt(0)
    for bagian in dok.sections:
        bagian.top_margin = Cm(2)
        bagian.bottom_margin = Cm(2)
        bagian.left_margin = Cm(2.5)
        bagian.right_margin = Cm(2)


def _p(dok, teks="", *, tebal=False, tengah=False, ukuran=None, sebelum=0, sesudah=0):
    p = dok.add_paragraph()
    if tengah:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(sebelum)
    p.paragraph_format.space_after = Pt(sesudah)
    if teks:
        r = p.add_run(teks)
        r.bold = tebal
        r.font.name = HURUF
        r.font.size = ukuran or UKURAN
    return p


def _sel(sel, teks, *, tebal=False, tengah=False):
    p = sel.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if tengah:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(teks)
    r.bold = tebal
    r.font.name = HURUF
    r.font.size = UKURAN


def _identitas(dok, baris):
    """Blok identitas tanpa garis — dua kolom, titik dua sejajar."""
    t = dok.add_table(rows=len(baris), cols=3)
    t.autofit = False
    for i, (label, isi) in enumerate(baris):
        lebar = (Cm(4.2), Cm(0.5), Cm(11.3))
        for j, teks in enumerate((label, ":", isi)):
            t.cell(i, j).width = lebar[j]
            _sel(t.cell(i, j), teks)
    return t


def _tanda_tangan(dok, kiri, kanan):
    """Dua kolom tanda tangan, tanpa garis. Ruang tanda tangannya tiga baris."""
    t = dok.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for sel, isi in zip(t.rows[0].cells, (kiri, kanan)):
        sel.width = Cm(7.8)
        p = sel.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for n, teks in enumerate(isi):
            if n:
                p = sel.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(teks)
            r.font.name = HURUF
            r.font.size = UKURAN
    return t


# ------------------------------------------------------------- kelengkapan
def bangun_kelengkapan(tujuan):
    """Formulir Daftar Kelengkapan Persyaratan — bentuknya mengikuti lampiran.

    Satu baris tabel bertanda `{{*butir.…}}`; docxgen mengulang baris itu
    sebanyak butir yang ada, jadi formulir jenis hak lain yang butirnya lebih
    banyak atau lebih sedikit tetap memakai template yang sama.
    """
    dok = Document()
    _atur_gaya(dok)

    _p(dok, "{{judul_formulir}}", tebal=True, tengah=True, ukuran=Pt(12))
    _p(dok, "{{dasar_formulir}}", tengah=True, ukuran=Pt(9), sesudah=10)

    _identitas(dok, [
        ("Nomor pradaftar", "{{nomor}}"),
        ("Tanggal", "{{tanggal}}"),
        ("Nama pemohon", "{{pemohon}}"),
        ("NIK", "{{pemohon_nik}}"),
        ("Alamat", "{{pemohon_alamat}}"),
        ("Letak tanah", "{{letak}}"),
        ("Luas menurut surat", "{{luas}}"),
    ])
    _p(dok, sesudah=8)

    t = dok.add_table(rows=2, cols=4)
    t.style = "Table Grid"
    t.autofit = False
    lebar = (Cm(1.3), Cm(10.5), Cm(1.9), Cm(2.1))

    kepala = ("No.", "Dokumen Persyaratan", "Ada", "Tidak Ada")
    for j, teks in enumerate(kepala):
        _sel(t.cell(0, j), teks, tebal=True, tengah=True)
        t.cell(0, j).width = lebar[j]
    # Baris contoh: seluruhnya diganti docxgen, satu baris per butir.
    isi = ("{{*butir.penanda}}", "{{*butir.nama}}",
           "{{*butir.ada}}", "{{*butir.tidak}}")
    for j, teks in enumerate(isi):
        _sel(t.cell(1, j), teks, tengah=(j != 1))
        t.cell(1, j).width = lebar[j]
    t.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    _p(dok, sesudah=6)
    _p(dok, "Keterangan:", tebal=True)
    _p(dok, "*) Coret yang tidak perlu")
    _p(dok, "{{?catatan}}Catatan: {{catatan}}", sebelum=4)

    _p(dok, sesudah=14)
    _tanda_tangan(dok, ["Pemohon,", "", "", "", "{{pemohon}}"],
                  ["{{kota}}, {{tanggal}}", "Petugas Loket,", "", "", "{{petugas}}"])
    dok.save(tujuan)
    return tujuan


# ------------------------------------------------------------ pengembalian
def bangun_pengembalian(tujuan):
    """Surat pengembalian berkas: hanya butir yang kurang atau perlu koreksi.

    Keduanya satu tabel (`dikembalikan`), dibedakan kolom Keterangan: "Belum
    ada" atau "Perlu diperbaiki: <catatan petugas>".
    """
    dok = Document()
    _atur_gaya(dok)

    _p(dok, "PENGEMBALIAN BERKAS PERMOHONAN", tebal=True, tengah=True, ukuran=Pt(12))
    _p(dok, "Nomor: {{nomor}}", tengah=True, sesudah=12)

    _p(dok, "Kepada Yth.")
    _p(dok, "Sdr. {{pemohon}}")
    _p(dok, "{{pemohon_alamat}}", sesudah=10)

    _p(dok, "Dengan hormat,", sesudah=6)
    _p(dok, "Setelah diperiksa terhadap {{judul_formulir}} "
            "({{dasar_formulir}}), berkas permohonan {{nama_hak}} atas nama "
            "{{pemohon}}{{letak_frasa}} belum dapat kami daftar karena "
            "kelengkapannya belum terpenuhi.", sesudah=8)
    _p(dok, "Dokumen yang masih perlu dilengkapi atau diperbaiki:", sesudah=4)

    t = dok.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    t.autofit = False
    lebar = (Cm(1.3), Cm(8.5), Cm(6.0))
    for j, teks in enumerate(("No.", "Dokumen", "Keterangan")):
        _sel(t.cell(0, j), teks, tebal=True, tengah=True)
        t.cell(0, j).width = lebar[j]
    isi = ("{{*dikembalikan._nomor}}", "{{*dikembalikan.nama}}",
           "{{*dikembalikan.keterangan}}")
    for j, teks in enumerate(isi):
        _sel(t.cell(1, j), teks, tengah=(j == 0))
        t.cell(1, j).width = lebar[j]

    _p(dok, sesudah=8)
    _p(dok, "Berkas dapat diajukan kembali setelah dokumen di atas dilengkapi "
            "atau diperbaiki. Atas perhatiannya kami ucapkan terima kasih.", sesudah=14)

    _tanda_tangan(dok, ["Diterima kembali oleh pemohon,", "", "", "", "{{pemohon}}"],
                  ["{{kota}}, {{tanggal}}", "Petugas Loket,", "", "", "{{petugas}}"])
    dok.save(tujuan)
    return tujuan


BENTUK = {"checklist.docx": bangun_kelengkapan,
          "pengembalian.docx": bangun_pengembalian}


def main():
    os.makedirs(jalur.TEMPLATE, exist_ok=True)
    for nama, bangun in BENTUK.items():
        tujuan = os.path.join(jalur.TEMPLATE, nama)
        bangun(tujuan)
        print(f"  dibuat  {tujuan}")


if __name__ == "__main__":
    main()
