# -*- coding: utf-8 -*-
"""Uji lapisan web Flask.

Yang dijaga di sini: kode jawaban, alamat tujuan, dan bentuk kuki sesi
untuk tiap rute. Uji inilah yang dipakai sebagai jaring pengaman saat rute
dipindahkan dari http.server ke Flask, dan tetap berlaku sesudahnya.

Basis datanya salinan sementara, sama seperti uji acuan dokumen.
"""
import io
import os
import re
import shutil
import sys
import tempfile
import unittest

from flask.testing import FlaskClient

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from berkas import db                                                 # noqa: E402
from berkas.dokumen import templat                                     # noqa: E402


POLA_CSRF = re.compile(r'name="_csrf" value="([^"]+)"')


class Klien(FlaskClient):
    """Klien uji yang berperilaku seperti peramban soal CSRF.

    Dua hal yang ditirukan: token dipungut dari halaman yang baru dibaca, dan
    tiap POST membawanya serta. Tanpa ini tiap uji harus mengurus tokennya
    sendiri, dan yang teruji jadi ketelitian ujinya, bukan aplikasinya.

    Uji yang memang hendak menembus penjagaan mengirim _csrf sendiri — yang
    dikirim eksplisit tidak pernah ditimpa.
    """

    # Halaman yang dijamin punya formulir POST untuk keduanya, admin maupun
    # petugas. Daftar berkas tidak bisa dipakai: di sana tak ada satu pun
    # <form method="post">, jadi tak ada token yang bisa dipungut.
    SUMBER_TOKEN = "/pengaturan"

    csrf = ""
    _sesi_token = None            # nilai kuki sesi saat token terakhir dipungut

    def _kuki_sesi(self):
        k = self.get_cookie("sesi")
        return k.value if k else None

    def segarkan(self):
        """Pungut token yang berlaku sekarang. Halaman masuk dipakai untuk yang
        belum masuk; kalau ia mengalihkan, berarti sudah masuk dan tokennya
        diambil dari Pengaturan."""
        j = self._pungut(super().open("/masuk"))
        if j.status_code != 200:
            self._pungut(super().open(self.SUMBER_TOKEN))

    def open(self, *a, **kw):
        if kw.get("method") == "POST":
            if not self.csrf or self._kuki_sesi() != self._sesi_token:
                # Klien baru, atau sesinya berganti: token lama sudah mati.
                self.segarkan()
            data = kw.get("data")
            if data is None:
                kw["data"] = {"_csrf": self.csrf}
            elif isinstance(data, dict) and "_csrf" not in data:
                kw["data"] = dict(data, _csrf=self.csrf)
        return self._pungut(super().open(*a, **kw))

    def _pungut(self, j):
        if j.mimetype == "text/html":
            m = POLA_CSRF.search(j.get_data(as_text=True))
            if m:
                self.csrf = m.group(1)
                self._sesi_token = self._kuki_sesi()
        return j


class Dasar(unittest.TestCase):
    """Tiap kelas uji mulai dari basis data kosong yang baru dibuat.

    Rangkaian uji yang sama dijalankan di dua basis data. Setel DB_URL ke
    PostgreSQL untuk menjalankannya di sana:

        DB_URL=postgresql://panitia:panitia@127.0.0.1:5432/panitia_a_uji             python -m unittest uji.uji_web

    Cara mengosongkannya berbeda: SQLite cukup berkas baru di folder
    sementara, PostgreSQL skemanya dibuang lalu dibuat ulang.
    """

    # Singgahan pratinjau ikut dialihkan ke folder sementara: cetakan loket
    # menaruh DOCX-nya di sana, dan uji tidak boleh menulis ke folder kerja
    # yang dipakai sungguhan.
    _dir_pratinjau = None

    @classmethod
    def _alihkan_pratinjau(cls):
        from berkas import pemeliharaan
        from berkas.dokumen import pdf
        cls._dir_pratinjau = tempfile.mkdtemp(prefix="uji-pratinjau-")
        cls._pratinjau_asli = (pdf.DIR_PRATINJAU, pemeliharaan.DIR_PRATINJAU)
        pdf.DIR_PRATINJAU = pemeliharaan.DIR_PRATINJAU = cls._dir_pratinjau

    @classmethod
    def _pulihkan_pratinjau(cls):
        if not cls._dir_pratinjau:
            return
        from berkas import pemeliharaan
        from berkas.dokumen import pdf
        pdf.DIR_PRATINJAU, pemeliharaan.DIR_PRATINJAU = cls._pratinjau_asli
        shutil.rmtree(cls._dir_pratinjau, ignore_errors=True)
        cls._dir_pratinjau = None

    @classmethod
    def setUpClass(cls):
        cls.dir = None
        cls._alihkan_pratinjau()
        if db.postgres():
            k = db.sambung()
            k.execute("DROP SCHEMA IF EXISTS public CASCADE")
            k.execute("CREATE SCHEMA public")
            k.commit()
            k.close()
            db.siapkan().close()
        else:
            cls.dir = tempfile.mkdtemp(prefix="uji-web-")
            cls._db_asli, cls._data_asli = db.BERKAS_DB, db.DIR_DATA
            db.DIR_DATA = cls.dir
            db.BERKAS_DB = os.path.join(cls.dir, "berkas.db")
            db.siapkan().close()

        from berkas.aplikasi import buat_aplikasi
        cls.app = buat_aplikasi()
        cls.app.config["TESTING"] = True
        cls.app.test_client_class = Klien

    @classmethod
    def tearDownClass(cls):
        cls._pulihkan_pratinjau()
        if cls.dir:
            db.BERKAS_DB, db.DIR_DATA = cls._db_asli, cls._data_asli
            shutil.rmtree(cls.dir, ignore_errors=True)

    def klien(self, masuk=False):
        c = self.app.test_client()
        c.get("/masuk")                       # memungut token tamu
        if masuk:
            j = c.post("/masuk", data={"username": "admin", "sandi": "admin123"})
            self.assertEqual(j.status_code, 303, "gagal masuk dengan pengguna bawaan")
            c.segarkan()                      # sesi baru, token baru
        return c


class Masuk(Dasar):
    def test_halaman_masuk_terbuka_tanpa_sesi(self):
        j = self.klien().get("/masuk")
        self.assertEqual(j.status_code, 200)
        self.assertIn("Berkas Panitia A", j.get_data(as_text=True))

    def test_sandi_salah_401_dan_tanpa_kuki(self):
        j = self.klien().post("/masuk", data={"username": "admin", "sandi": "salah"})
        self.assertEqual(j.status_code, 401)
        self.assertNotIn("sesi=", j.headers.get("Set-Cookie", ""))

    def test_pengguna_tak_dikenal_401(self):
        j = self.klien().post("/masuk", data={"username": "hantu", "sandi": "apa saja"})
        self.assertEqual(j.status_code, 401)

    def test_kepala_nosniff_ikut_di_rute_flask(self):
        # _kirim() yang lama memasangnya di tiap jawaban; jangan sampai hilang
        # justru di halaman masuk.
        self.assertEqual(self.klien().get("/masuk").headers["X-Content-Type-Options"],
                         "nosniff")

    def test_berhasil_masuk_kuki_httponly_samesite(self):
        j = self.klien().post("/masuk", data={"username": "admin", "sandi": "admin123"})
        self.assertEqual(j.status_code, 303)
        self.assertEqual(j.headers["Location"], "/berkas")
        kuki = j.headers["Set-Cookie"]
        self.assertIn("HttpOnly", kuki)
        self.assertIn("SameSite=Lax", kuki)
        self.assertIn("Path=/", kuki)

    def test_sudah_masuk_tidak_lihat_halaman_masuk_lagi(self):
        j = self.klien(masuk=True).get("/masuk")
        self.assertEqual(j.status_code, 303)
        self.assertEqual(j.headers["Location"], "/berkas")

    def test_keluar_menghapus_sesi(self):
        c = self.klien(masuk=True)
        j = c.get("/keluar")
        self.assertEqual(j.status_code, 303)
        self.assertEqual(j.headers["Location"], "/masuk")
        self.assertEqual(c.get("/berkas").headers["Location"], "/masuk")

    def test_token_tetap_sah_lintas_permintaan(self):
        c = self.klien(masuk=True)
        for _ in range(3):
            self.assertEqual(c.get("/berkas").status_code, 200)


class Penjagaan(Dasar):
    def test_tanpa_sesi_dialihkan_ke_masuk(self):
        c = self.klien()
        for jalur in ("/", "/berkas", "/berkas/baru", "/pradaftar", "/pradaftar/baru",
                      "/referensi", "/template", "/pengaturan"):
            j = c.get(jalur)
            self.assertEqual(j.status_code, 303, jalur)
            self.assertEqual(j.headers["Location"], "/masuk", jalur)

    def test_kartu_lipat_tanpa_sesi_dijawab_401_bukan_halaman_masuk(self):
        # app.js menelan jawabannya mentah-mentah ke dalam kartu; halaman
        # masuk di sana akan tampak seperti kartu yang rusak.
        j = self.klien().get("/referensi/bagian/kecamatan")
        self.assertEqual(j.status_code, 401)
        self.assertEqual(j.get_data(as_text=True), "sesi habis")

    def test_post_tanpa_sesi_dialihkan(self):
        j = self.klien().post("/berkas/baru", data={"status": "draf"})
        self.assertEqual(j.status_code, 303)
        self.assertEqual(j.headers["Location"], "/masuk")

    def test_kuki_palsu_diperlakukan_seperti_tanpa_sesi(self):
        c = self.klien()
        c.set_cookie("sesi", "token-karangan", domain="localhost")
        self.assertEqual(c.get("/berkas").headers["Location"], "/masuk")


class Pengguna(Dasar):
    def test_ganti_sandi_lama_salah_ditolak(self):
        j = self.klien(masuk=True).post("/sandi", data={"lama": "bukan", "baru": "baru123"})
        self.assertEqual(j.headers["Location"], "/pengaturan?pesan=sandi_salah#p-pengguna")

    def test_ganti_sandi_kosong_ditolak(self):
        j = self.klien(masuk=True).post("/sandi", data={"lama": "admin123", "baru": ""})
        self.assertEqual(j.headers["Location"], "/pengaturan?pesan=sandi_salah#p-pengguna")

    def test_tambah_pengguna_lalu_bisa_masuk(self):
        c = self.klien(masuk=True)
        j = c.post("/pengguna", data={"nama": "Uji Coba", "username": "ujicoba",
                                      "sandi": "rahasia123"})
        self.assertEqual(j.headers["Location"], "/pengaturan?pesan=pengguna#p-pengguna")
        j = self.app.test_client().post(
            "/masuk", data={"username": "ujicoba", "sandi": "rahasia123"})
        self.assertEqual(j.status_code, 303)

    def test_username_kembar_ditolak(self):
        c = self.klien(masuk=True)
        c.post("/pengguna", data={"nama": "Satu", "username": "kembar", "sandi": "a"})
        j = c.post("/pengguna", data={"nama": "Dua", "username": "kembar", "sandi": "b"})
        self.assertEqual(j.headers["Location"], "/pengaturan?pesan=pengguna_ada#p-pengguna")

    def test_isian_kurang_diabaikan(self):
        j = self.klien(masuk=True).post("/pengguna", data={"nama": "", "username": "x",
                                                           "sandi": "y"})
        self.assertEqual(j.headers["Location"], "/pengaturan")


class Berkas(Dasar):
    """Blueprint rute/berkas.py — daftar, formulir, dan pencetakan."""

    def test_tiap_halaman_terakit_penuh(self):
        c = self.klien(masuk=True)
        for jalur in ("/", "/berkas", "/berkas/baru", "/pradaftar", "/pradaftar/baru",
                      "/referensi", "/template", "/pengaturan"):
            j = c.get(jalur)
            self.assertEqual(j.status_code, 200, jalur)
            self.assertEqual(j.headers["Content-Type"], "text/html; charset=utf-8", jalur)
            self.assertIn("<!doctype html", j.get_data(as_text=True).lower(), jalur)

    def test_kepala_nosniff_ikut_terbawa(self):
        j = self.klien(masuk=True).get("/berkas")
        self.assertEqual(j.headers["X-Content-Type-Options"], "nosniff")

    def test_kartu_lipat_ditandai_supaya_app_js_mengenalinya(self):
        j = self.klien(masuk=True).get("/referensi/bagian/kecamatan")
        self.assertEqual(j.status_code, 200)
        self.assertEqual(j.headers.get("X-Potongan"), "1")

    def test_kabar_dari_query_tampil(self):
        j = self.klien(masuk=True).get("/berkas?pesan=tersimpan")
        self.assertIn("Tersimpan", j.get_data(as_text=True))

    def test_daftar_berkas_disaring_dan_dipaging(self):
        c = self.klien(masuk=True)
        for _ in range(3):
            c.post("/berkas/baru", data={"status": "draf", "jenis_kegiatan": "baru",
                                         "jenis_hak_dimohon": "HM"})
        c.post("/berkas/baru", data={"status": "selesai", "jenis_kegiatan": "baru",
                                     "jenis_hak_dimohon": "WAKAF"})
        isi = c.get("/berkas").get_data(as_text=True)
        self.assertIn('name="hak"', isi)
        self.assertIn('<option value="WAKAF"', isi)

        isi = c.get("/berkas?hak=WAKAF").get_data(as_text=True)
        self.assertEqual(isi.count('<td class="nomor">'), 3)       # no, risalah, SK
        self.assertIn("saring-cap", isi)
        isi = c.get("/berkas?hak=HM&status=selesai").get_data(as_text=True)
        self.assertIn("Tidak ada berkas yang cocok", isi)

        isi = c.get("/berkas?per=10&hal=99").get_data(as_text=True)
        self.assertEqual(c.get("/berkas?per=abc").status_code, 200)
        self.assertIn('aria-current="page"', isi)                  # hal dijepit ke akhir

    def test_ctrl_s_dipasang_pada_formulir_yang_boleh_diubah(self):
        c = self.klien(masuk=True)
        for jalur in ("/berkas/baru", "/pradaftar/baru"):
            self.assertIn("data-simpan-pintas", c.get(jalur).get_data(as_text=True), jalur)
        j = c.post("/berkas/baru", data={"status": "draf", "jenis_kegiatan": "baru"})
        isi = c.get(j.headers["Location"].split("?")[0]).get_data(as_text=True)
        self.assertIn("data-simpan-pintas", isi)
        self.assertIn('class="pintas"', isi)

    def test_nomor_halaman_berelipsis(self):
        from berkas import web
        self.assertEqual(web._nomor_halaman(1, 3), [1, 2, 3])
        self.assertEqual(web._nomor_halaman(6, 12), [1, 2, None, 5, 6, 7, None, 11, 12])
        self.assertEqual(web._nomor_halaman(4, 12), [1, 2, 3, 4, 5, None, 11, 12])

    def test_jalur_tak_dikenal_404_berhalaman(self):
        j = self.klien(masuk=True).get("/entah-apa")
        self.assertEqual(j.status_code, 404)
        self.assertIn("tidak ditemukan", j.get_data(as_text=True).lower())

    def test_buat_lalu_hapus_berkas(self):
        c = self.klien(masuk=True)
        j = c.post("/berkas/baru", data={"status": "draf", "jenis_kegiatan": "baru"})
        self.assertEqual(j.status_code, 303)
        tujuan = j.headers["Location"]
        self.assertIn("/berkas/", tujuan)
        bid = tujuan.split("/berkas/")[1].split("?")[0]
        self.assertEqual(c.get(f"/berkas/{bid}").status_code, 200)
        j = c.post(f"/berkas/{bid}/hapus")
        self.assertEqual(j.headers["Location"], "/berkas?pesan=dihapus")


class Pengaturan(Dasar):
    """Blueprint rute/pengaturan.py."""

    def test_isian_berawalan_set_tersimpan(self):
        c = self.klien(masuk=True)
        j = c.post("/pengaturan", data={"set_nama_kantor": "Kantah Uji Bone Bolango",
                                        "bukan_setelan": "diabaikan"})
        self.assertEqual(j.headers["Location"], "/pengaturan?pesan=pengaturan#p-kantor")
        self.assertIn("Kantah Uji Bone Bolango",
                      c.get("/pengaturan").get_data(as_text=True))

    def test_yang_tanpa_awalan_set_tidak_ikut_masuk(self):
        c = self.klien(masuk=True)
        c.post("/pengaturan", data={"bukan_setelan": "jangan-disimpan"})
        k = db.sambung()
        try:
            ada = k.execute("SELECT 1 FROM pengaturan WHERE kunci=?",
                            ("bukan_setelan",)).fetchone()
        finally:
            k.close()
        self.assertIsNone(ada)

    def test_tiap_aksi_pemeliharaan_dijawab(self):
        c = self.klien(masuk=True)
        for aksi in ("pratinjau", "cadangan_db", "keluaran"):
            j = c.post("/pemeliharaan", data={"aksi": aksi, "umur_hari": "90"})
            self.assertEqual(j.status_code, 303, aksi)
            self.assertIn("#penyimpanan", j.headers["Location"], aksi)

    def test_umur_hari_bukan_angka_tidak_meledak(self):
        j = self.klien(masuk=True).post("/pemeliharaan",
                                        data={"aksi": "keluaran", "umur_hari": "entah"})
        self.assertEqual(j.status_code, 303)

    def test_aksi_tak_dikenal_dianggap_tidak_ada_yang_dibuang(self):
        j = self.klien(masuk=True).post("/pemeliharaan", data={"aksi": "ngawur"})
        self.assertIn("pesan=bersih_kosong", j.headers["Location"])


class Unduhan(Dasar):
    """Blueprint rute/unduhan.py."""

    def test_dokumen_tak_ada_404(self):
        j = self.klien(masuk=True).get("/unduh/tidak-ada.docx")
        self.assertEqual(j.status_code, 404)
        self.assertIn("text/plain", j.headers["Content-Type"])

    def test_foto_tak_ada_404(self):
        self.assertEqual(self.klien(masuk=True).get("/foto/999999").status_code, 404)

    def test_pdf_dokumen_tak_ada_404(self):
        self.assertEqual(
            self.klien(masuk=True).get("/pdf/tidak-ada.docx").status_code, 404)

    def test_pratinjau_dokumen_tak_ada_berhalaman_bukan_teks(self):
        # Halaman ini dibuka dari tab Cetak; galatnya harus tetap berbentuk
        # halaman supaya petugas bisa kembali ke berkasnya.
        j = self.klien(masuk=True).get("/pratinjau/tidak-ada.docx")
        self.assertEqual(j.status_code, 404)
        self.assertIn("<!doctype html", j.get_data(as_text=True).lower())
        self.assertIn("belum pernah dicetak", j.get_data(as_text=True))

    def test_nama_berjalur_dipangkas_ke_nama_berkas_saja(self):
        # /unduh/../../data/berkas.db tidak boleh keluar dari folder keluaran
        j = self.klien(masuk=True).get("/unduh/..%2F..%2Fdata%2Fberkas.db")
        self.assertEqual(j.status_code, 404)

    def test_semua_perlu_masuk(self):
        c = self.klien()
        for jalur in ("/unduh/a.docx", "/foto/1", "/pdf/a.docx", "/pratinjau/a.docx"):
            self.assertEqual(c.get(jalur).headers.get("Location"), "/masuk", jalur)


class Cetak(Dasar):
    """Pencetakan hanya boleh jalan kalau pemeriksaannya bersih."""

    def test_berkas_kosong_ditolak_cetak(self):
        c = self.klien(masuk=True)
        j = c.post("/berkas/baru", data={"status": "draf", "jenis_kegiatan": "baru"})
        bid = j.headers["Location"].split("/berkas/")[1].split("?")[0]
        j = c.post(f"/berkas/{bid}/cetak", data={"jenis": "semua"})
        self.assertEqual(j.status_code, 303)
        self.assertEqual(j.headers["Location"], f"/berkas/{bid}?pesan=galat_cetak")

    def test_batas_unggahan_terpasang(self):
        self.assertEqual(self.app.config["MAX_CONTENT_LENGTH"], 120 * 1024 * 1024)


class Rbac(Dasar):
    """Dua peran: admin dan petugas.

    Yang diuji di sini penolakan di sisi server, bukan tombol yang hilang.
    Menyembunyikan tombol tidak menjaga apa pun — siapa saja bisa mengetik
    alamatnya langsung — jadi tiap larangan diuji dengan menembak rutenya.
    """

    def petugas(self, username, sandi="rahasia123"):
        """Buat satu petugas lewat admin, lalu kembalikan klien yang sudah masuk."""
        a = self.klien(masuk=True)
        a.post("/pengguna", data={"nama": username.title(), "username": username,
                                  "sandi": sandi, "peran": "petugas"})
        c = self.app.test_client()
        c.get("/masuk")
        j = c.post("/masuk", data={"username": username, "sandi": sandi})
        self.assertEqual(j.status_code, 303, f"petugas {username} gagal masuk")
        c.segarkan()
        return c

    def buat_berkas(self, c):
        j = c.post("/berkas/baru", data={"status": "draf", "jenis_kegiatan": "baru",
                                         "penerima_nama": "Uji Penerima"})
        self.assertEqual(j.status_code, 303)
        return int(j.headers["Location"].split("/berkas/")[1].split("?")[0])

    def berkas_yatim(self):
        """Berkas tanpa pemilik, seperti hasil impor Excel."""
        k = db.sambung()
        try:
            bid = db.sisip_id(k, "INSERT INTO berkas (jenis_hak_dimohon,dibuat_oleh) "
                                 "VALUES ('HM', NULL)")
            k.commit()
        finally:
            k.close()
        return bid

    # ----- yang boleh
    def test_petugas_bisa_menambah_dan_mengubah_berkasnya(self):
        c = self.petugas("tugas1")
        bid = self.buat_berkas(c)
        self.assertEqual(c.get(f"/berkas/{bid}").status_code, 200)
        j = c.post(f"/berkas/{bid}", data={"status": "draf", "penerima_nama": "Diubah"})
        self.assertEqual(j.status_code, 303)
        self.assertIn("Diubah", c.get(f"/berkas/{bid}").get_data(as_text=True))

    def test_petugas_bisa_mencetak_berkasnya_sendiri(self):
        # Berkasnya sengaja belum lengkap, jadi berhenti di galat_cetak —
        # yang diuji izinnya, bukan perakitan dokumennya.
        c = self.petugas("tugas2")
        bid = self.buat_berkas(c)
        j = c.post(f"/berkas/{bid}/cetak", data={"jenis": "semua"})
        self.assertEqual(j.status_code, 303)
        self.assertEqual(j.headers["Location"], f"/berkas/{bid}?pesan=galat_cetak")

    def test_petugas_boleh_membaca_data_referensi(self):
        c = self.petugas("tugas3")
        self.assertEqual(c.get("/referensi").status_code, 200)
        self.assertEqual(c.get("/referensi/bagian/kecamatan").status_code, 200)

    def test_petugas_boleh_mengganti_sandinya_sendiri(self):
        c = self.petugas("tugas4")
        j = c.post("/sandi", data={"lama": "rahasia123", "baru": "sandibaru9"})
        self.assertEqual(j.headers["Location"], "/pengaturan?pesan=sandi#p-pengguna")

    # ----- yang tidak boleh
    def test_petugas_tidak_bisa_menghapus_berkas(self):
        c = self.petugas("tugas5")
        bid = self.buat_berkas(c)
        j = c.post(f"/berkas/{bid}/hapus")
        self.assertEqual(j.status_code, 303)
        self.assertIn("galat=", j.headers["Location"])
        self.assertEqual(c.get(f"/berkas/{bid}").status_code, 200, "berkasnya ikut hilang")

    def test_petugas_ditolak_di_semua_post_referensi(self):
        c = self.petugas("tugas6")
        jalur = ["/referensi/kecamatan", "/referensi/kecamatan/hapus",
                 "/referensi/desa", "/referensi/desa/hapus",
                 "/referensi/desa/pejabat", "/referensi/desa/pejabat/aktif",
                 "/referensi/desa/pejabat/hapus", "/referensi/wilayah",
                 "/referensi/panitia", "/referensi/panitia/hapus",
                 "/referensi/panitia/salin", "/referensi/panitia/anggota",
                 "/referensi/panitia/anggota/hapus"]
        for j in jalur:
            r = c.post(j, data={"nama": "Sisipan", "id": "1"})
            self.assertEqual(r.status_code, 303, j)
            self.assertIn("galat=", r.headers["Location"], j)
        k = db.sambung()
        try:
            ada = k.execute("SELECT 1 FROM ref_kecamatan WHERE nama=?",
                            ("Sisipan",)).fetchone()
        finally:
            k.close()
        self.assertIsNone(ada, "ada POST referensi yang lolos")

    def test_petugas_ditolak_di_seluruh_halaman_template(self):
        c = self.petugas("tugas7")
        self.assertEqual(c.get("/template").status_code, 403)
        self.assertEqual(c.get("/template/bagian/cara").status_code, 403)
        self.assertEqual(c.get("/template/unduh/sk").status_code, 403)
        self.assertEqual(c.post("/template/unggah", data={"jenis": "sk"}).status_code, 303)

    def test_petugas_ditolak_mengubah_pengaturan_dan_akun(self):
        c = self.petugas("tugas8")
        for jalur, data in (("/pengaturan", {"set_nama_kantor": "Dibajak"}),
                            ("/pemeliharaan", {"aksi": "pratinjau"}),
                            ("/pengguna", {"nama": "X", "username": "x", "sandi": "y"}),
                            ("/pengguna/peran", {"id": "1", "peran": "petugas"})):
            r = c.post(jalur, data=data)
            self.assertEqual(r.status_code, 303, jalur)
            self.assertIn("galat=", r.headers["Location"], jalur)
        k = db.sambung()
        try:
            self.assertIsNone(k.execute("SELECT 1 FROM pengguna WHERE username='x'").fetchone())
            self.assertEqual(k.execute("SELECT peran FROM pengguna WHERE id=1")
                             .fetchone()["peran"], "admin", "admin berhasil diturunkan")
        finally:
            k.close()

    # ----- batas antarpetugas
    def test_petugas_tidak_melihat_berkas_petugas_lain(self):
        satu = self.petugas("alfa")
        dua = self.petugas("beta")
        bid = self.buat_berkas(satu)
        self.assertEqual(dua.get(f"/berkas/{bid}").status_code, 403)
        self.assertNotIn(f'/berkas/{bid}"', dua.get("/berkas").get_data(as_text=True))

    def test_petugas_tidak_bisa_mengubah_berkas_petugas_lain(self):
        satu = self.petugas("gama")
        dua = self.petugas("delta")
        bid = self.buat_berkas(satu)
        self.assertEqual(dua.post(f"/berkas/{bid}", data={"status": "selesai"}).status_code,
                         303)
        k = db.sambung()
        try:
            status = k.execute("SELECT status FROM berkas WHERE id=?", (bid,)).fetchone()
        finally:
            k.close()
        self.assertEqual(status["status"], "draf", "berkas orang lain ikut berubah")

    def test_petugas_tidak_bisa_mencetak_berkas_orang_lain(self):
        satu = self.petugas("epsilon")
        dua = self.petugas("zeta")
        bid = self.buat_berkas(satu)
        j = dua.post(f"/berkas/{bid}/cetak", data={"jenis": "semua"})
        self.assertIn("galat=", j.headers["Location"])

    # ----- arsip tanpa pemilik
    def test_arsip_tanpa_pemilik_terlihat_petugas(self):
        c = self.petugas("eta")
        bid = self.berkas_yatim()
        self.assertEqual(c.get(f"/berkas/{bid}").status_code, 200)
        self.assertIn(f'/berkas/{bid}"', c.get("/berkas").get_data(as_text=True))

    def test_arsip_tanpa_pemilik_baca_saja_bagi_petugas(self):
        c = self.petugas("theta")
        bid = self.berkas_yatim()
        isi = c.get(f"/berkas/{bid}").get_data(as_text=True)
        self.assertIn("baca-saja", isi, "isian mestinya dimatikan")
        self.assertNotIn("data-simpan-pintas", isi, "Ctrl+S tak boleh aktif di baca-saja")
        self.assertNotIn('type="submit" form="form-berkas">Simpan', isi)

    def test_petugas_tidak_bisa_menulis_ke_arsip_tanpa_pemilik(self):
        c = self.petugas("iota")
        bid = self.berkas_yatim()
        self.assertEqual(c.post(f"/berkas/{bid}", data={"status": "selesai"}).status_code,
                         303)
        k = db.sambung()
        try:
            r = k.execute("SELECT status, dibuat_oleh FROM berkas WHERE id=?",
                          (bid,)).fetchone()
        finally:
            k.close()
        self.assertEqual(r["status"], "draf")
        self.assertIsNone(r["dibuat_oleh"], "kepemilikan ikut berpindah diam-diam")

    def test_admin_boleh_menulis_ke_arsip_tanpa_pemilik(self):
        bid = self.berkas_yatim()
        a = self.klien(masuk=True)
        self.assertEqual(a.get(f"/berkas/{bid}").status_code, 200)
        self.assertEqual(a.post(f"/berkas/{bid}", data={"status": "diperiksa"}).status_code,
                         303)

    # ----- unduhan ikut aturan kepemilikan
    def test_dokumen_tanpa_catatan_terbit_hanya_untuk_admin(self):
        c = self.petugas("kappa")
        self.assertEqual(c.get("/unduh/entah.docx").status_code, 403)
        self.assertEqual(c.get("/pdf/entah.docx").status_code, 403)
        self.assertEqual(self.klien(masuk=True).get("/unduh/entah.docx").status_code, 404)

    def test_foto_berkas_orang_lain_ditolak(self):
        satu = self.petugas("lambda1")
        dua = self.petugas("mu")
        bid = self.buat_berkas(satu)
        k = db.sambung()
        try:
            fid = db.sisip_id(k, "INSERT INTO foto_lapang (berkas_id,nama_file,urut) "
                                 "VALUES (?,?,0)", (bid, "tidak-ada.jpg"))
            k.commit()
        finally:
            k.close()
        self.assertEqual(dua.get(f"/foto/{fid}").status_code, 403)
        self.assertEqual(satu.get(f"/foto/{fid}").status_code, 404)   # datanya ada, filenya tidak

    # ----- pengelolaan peran
    def test_admin_bisa_menaikkan_dan_menurunkan_peran(self):
        c = self.petugas("nu")
        k = db.sambung()
        try:
            uid = k.execute("SELECT id FROM pengguna WHERE username='nu'").fetchone()["id"]
        finally:
            k.close()
        a = self.klien(masuk=True)
        j = a.post("/pengguna/peran", data={"id": uid, "peran": "admin"})
        self.assertIn("pesan=peran", j.headers["Location"])
        self.assertEqual(c.get("/template").status_code, 200, "kenaikan peran belum berlaku")
        a.post("/pengguna/peran", data={"id": uid, "peran": "petugas"})
        self.assertEqual(c.get("/template").status_code, 403)

    def test_admin_tidak_bisa_menurunkan_dirinya_sendiri(self):
        a = self.klien(masuk=True)
        j = a.post("/pengguna/peran", data={"id": "1", "peran": "petugas"})
        self.assertIn("pesan=peran_sendiri", j.headers["Location"])
        self.assertEqual(a.get("/template").status_code, 200)

    def test_peran_karangan_jatuh_ke_petugas(self):
        a = self.klien(masuk=True)
        a.post("/pengguna", data={"nama": "Nekat", "username": "nekat",
                                  "sandi": "rahasia123", "peran": "superadmin"})
        k = db.sambung()
        try:
            peran = k.execute("SELECT peran FROM pengguna WHERE username='nekat'") \
                     .fetchone()["peran"]
        finally:
            k.close()
        self.assertEqual(peran, "petugas")

    def test_peran_tak_dikenal_pada_ubah_ditolak(self):
        a = self.klien(masuk=True)
        j = a.post("/pengguna/peran", data={"id": "2", "peran": "dewa"})
        self.assertIn("pesan=peran_salah", j.headers["Location"])

    # ----- tampilan
    def test_menu_template_hanya_untuk_admin(self):
        c = self.petugas("omikron")
        self.assertNotIn(">Template<", c.get("/berkas").get_data(as_text=True))
        self.assertIn(">Template<", self.klien(masuk=True).get("/berkas")
                      .get_data(as_text=True))

    def test_petugas_tidak_melihat_daftar_akun_di_pengaturan(self):
        c = self.petugas("pi")
        self.petugas("rho_tetangga")        # akun lain yang tidak boleh bocor
        j = c.get("/pengaturan")
        self.assertEqual(j.status_code, 200)
        isi = j.get_data(as_text=True)
        self.assertIn('action="/sandi"', isi, "petugas tetap bisa ganti sandi")
        self.assertNotIn('action="/pengguna"', isi, "formulir tambah pengguna bocor")
        self.assertNotIn("rho_tetangga", isi, "username akun lain bocor")
        self.assertNotIn("/pengguna/peran", isi, "pengubah peran bocor")

    def test_petugas_tidak_melihat_identitas_kantor(self):
        a = self.klien(masuk=True)
        a.post("/pengaturan", data={"set_nama_kantor": "Kantah Rahasia"})
        c = self.petugas("sigma")
        self.assertNotIn("Kantah Rahasia", c.get("/pengaturan").get_data(as_text=True))


class Tampilan(Dasar):
    """Kerangka halaman: menu samping dan pemecahan isi jadi tab."""

    def _panel(self, isi):
        """id tiap <div class="panel"> pada halaman."""
        return re.findall(r'<div class="panel" id="([^"]+)"', isi)

    def _tombol_tab(self, isi):
        """panel yang dituju tiap tombol tab."""
        return re.findall(r'<button type="button" role="tab"[^>]*data-panel="([^"]+)"', isi)

    def test_menu_samping_ada_di_tiap_halaman(self):
        c = self.klien(masuk=True)
        for jalur in ("/berkas", "/referensi", "/template", "/pengaturan"):
            isi = c.get(jalur).get_data(as_text=True)
            self.assertIn('<aside class="samping"', isi, jalur)
            self.assertIn('id="kuncup"', isi, jalur)       # tombol kuncup
            self.assertIn('id="buka-nav"', isi, jalur)     # tombol laci layar sempit
            self.assertIn('class="sprite"', isi, jalur)    # ikon

    def test_halaman_masuk_tanpa_menu_samping(self):
        isi = self.klien().get("/masuk").get_data(as_text=True)
        self.assertNotIn('class="samping"', isi)
        self.assertNotIn("buka-nav", isi)

    def test_menu_menandai_halaman_yang_sedang_dibuka(self):
        c = self.klien(masuk=True)
        for jalur in ("/berkas", "/referensi", "/template", "/pengaturan"):
            isi = c.get(jalur).get_data(as_text=True)
            aktif = re.findall(r'<a href="([^"]+)" class="samping-butir aktif"', isi)
            self.assertEqual(aktif, [jalur], jalur)

    def test_keadaan_kuncup_dipasang_sebelum_halaman_digambar(self):
        # Kalau skripnya di akhir badan, menunya sempat berkedip lebar dulu.
        isi = self.klien(masuk=True).get("/berkas").get_data(as_text=True)
        kepala = isi.split("</head>")[0]
        self.assertIn("menu-kuncup", kepala)

    def test_referensi_dipecah_lima_tab(self):
        isi = self.klien(masuk=True).get("/referensi").get_data(as_text=True)
        urut = ["r-hak", "r-klausa", "r-lengkap", "r-wilayah", "r-panitia"]
        self.assertEqual(self._tombol_tab(isi), urut)
        self.assertEqual(self._panel(isi), urut)

    def test_template_dipecah_lima_tab(self):
        isi = self.klien(masuk=True).get("/template").get_data(as_text=True)
        urut = ["t-hm", "t-wakaf", "t-hp", "t-loket", "t-penanda"]
        self.assertEqual(self._tombol_tab(isi), urut)
        self.assertEqual(self._panel(isi), urut)

    def test_template_terpisah_menurut_tata_naskahnya(self):
        """Cetakan loket berdiri sendiri: formulirnya berlaku untuk semua jenis
        hak, jadi tidak boleh menumpang di tab tata naskah mana pun."""
        isi = self.klien(masuk=True).get("/template").get_data(as_text=True)
        hm = isi.split('id="t-hm"')[1].split('id="t-wakaf"')[0]
        wakaf = isi.split('id="t-wakaf"')[1].split('id="t-hp"')[0]
        hp = isi.split('id="t-hp"')[1].split('id="t-loket"')[0]
        loket = isi.split('id="t-loket"')[1].split('id="t-penanda"')[0]
        self.assertEqual(sorted(re.findall(r'id="b-(tpl-[\w-]+)"', hp)),
                         ["tpl-bap-hp", "tpl-risalah-hp", "tpl-sk-hp"])
        self.assertEqual(sorted(re.findall(r'id="b-(tpl-[\w-]+)"', loket)),
                         ["tpl-checklist", "tpl-pengembalian"])
        self.assertEqual(sorted(re.findall(r'id="b-(tpl-[\w-]+)"', hm)),
                         ["tpl-bap", "tpl-risalah", "tpl-sk"])
        self.assertEqual(sorted(re.findall(r'id="b-(tpl-[\w-]+)"', wakaf)),
                         ["tpl-bap-wakaf", "tpl-risalah-wakaf", "tpl-sk-wakaf"])

    def test_tiap_tombol_tab_punya_panelnya(self):
        """Tombol tab yang menunjuk panel tak ada bikin halaman kosong diam-diam."""
        c = self.klien(masuk=True)
        for jalur in ("/referensi", "/template", "/pengaturan", "/berkas/baru"):
            isi = c.get(jalur).get_data(as_text=True)
            self.assertEqual(self._tombol_tab(isi), self._panel(isi), jalur)

    def test_label_tab_tidak_terkode_ganda(self):
        # _tab() sudah meng-escape; pemanggilnya harus mengirim '&' apa adanya.
        isi = self.klien(masuk=True).get("/template").get_data(as_text=True)
        self.assertNotIn("&amp;amp;", isi)

    def test_alihan_lama_masih_menunjuk_ke_dalam_tab(self):
        """POST referensi mengalihkan ke ?buka=<kunci>#b-<kunci>. Jangkar itu
        harus tetap berada di dalam salah satu panel, supaya app.js membuka
        tab yang benar."""
        isi = self.klien(masuk=True).get("/referensi?buka=desa").get_data(as_text=True)
        wilayah = isi.split('id="r-wilayah"')[1].split('<div class="panel"')[0]
        for jangkar in ("b-kemendagri", "b-kecamatan", "b-desa"):
            self.assertIn(f'id="{jangkar}"', wilayah, jangkar)
        panitia = isi.split('id="r-panitia"')[1]
        self.assertIn('id="daftar-sk"', panitia)

    def test_petugas_tidak_dapat_butir_menu_template(self):
        c = self.klien(masuk=True)
        c.post("/pengguna", data={"nama": "Tampil", "username": "tampil",
                                  "sandi": "rahasia123", "peran": "petugas"})
        p = self.app.test_client()
        p.post("/masuk", data={"username": "tampil", "sandi": "rahasia123"})
        isi = p.get("/berkas").get_data(as_text=True)
        self.assertNotIn('href="/template"', isi)
        self.assertIn('href="/referensi"', isi)


class Csrf(Dasar):
    """Token anti-CSRF.

    Klien uji di berkas ini menyisipkan token sendiri, seperti peramban. Uji
    di kelas ini mengirim _csrf secara eksplisit supaya yang diuji benar-benar
    penolakannya, bukan ketelitian kliennya.
    """

    POST_CONTOH = "/referensi/kecamatan"

    def test_tanpa_token_ditolak(self):
        j = self.klien(masuk=True).post(self.POST_CONTOH,
                                        data={"nama": "Tanpa Token", "_csrf": ""})
        self.assertEqual(j.status_code, 400)
        self.assertNotIn("Tanpa Token", self.klien(masuk=True)
                         .get("/referensi/bagian/kecamatan").get_data(as_text=True))

    def test_token_karangan_ditolak(self):
        j = self.klien(masuk=True).post(self.POST_CONTOH,
                                        data={"nama": "Karangan", "_csrf": "a" * 43})
        self.assertEqual(j.status_code, 400)

    def test_token_milik_sesi_lain_ditolak(self):
        satu, dua = self.klien(masuk=True), self.klien(masuk=True)
        self.assertNotEqual(satu.csrf, dua.csrf, "dua sesi mestinya beda token")
        j = dua.post(self.POST_CONTOH, data={"nama": "Pinjam", "_csrf": satu.csrf})
        self.assertEqual(j.status_code, 400)

    def test_token_sendiri_diterima(self):
        c = self.klien(masuk=True)
        j = c.post(self.POST_CONTOH, data={"nama": "Sah", "_csrf": c.csrf})
        self.assertEqual(j.status_code, 303)

    def test_lewat_kepala_permintaan_juga_diterima(self):
        c = self.klien(masuk=True)
        j = c.post(self.POST_CONTOH, data={"nama": "Lewat Kepala", "_csrf": ""},
                   headers={"X-CSRF-Token": c.csrf})
        self.assertEqual(j.status_code, 303)

    def test_get_tidak_pernah_dihalangi(self):
        c = self.klien(masuk=True)
        for jalur in ("/berkas", "/referensi", "/template", "/pengaturan",
                      "/referensi/bagian/kecamatan"):
            self.assertNotEqual(c.get(jalur).status_code, 400, jalur)

    def test_halaman_masuk_terjaga_kuki_tamu(self):
        # Tanpa ini penyerang bisa memaksa korban masuk ke akun miliknya.
        c = self.app.test_client()
        j = c.post("/masuk", data={"username": "admin", "sandi": "admin123",
                                   "_csrf": "bukan-token"})
        self.assertEqual(j.status_code, 400)

    def test_masuk_dengan_token_dari_halamannya_berhasil(self):
        c = self.app.test_client()
        isi = c.get("/masuk").get_data(as_text=True)
        tok = POLA_CSRF.search(isi).group(1)
        self.assertIn("csrf_tamu", str(c.get_cookie("csrf_tamu")))
        j = c.post("/masuk", data={"username": "admin", "sandi": "admin123", "_csrf": tok})
        self.assertEqual(j.status_code, 303)

    def test_token_berganti_sesudah_masuk_ulang(self):
        c = self.klien(masuk=True)
        lama = c.csrf
        c.get("/keluar")
        c.get("/masuk")
        c.post("/masuk", data={"username": "admin", "sandi": "admin123"})
        c.get("/pengaturan")
        self.assertNotEqual(c.csrf, lama, "sesi baru mestinya bertoken baru")

    def test_token_lama_mati_sesudah_keluar(self):
        c = self.klien(masuk=True)
        lama = c.csrf
        c.get("/keluar")
        d = self.klien(masuk=True)
        j = d.post(self.POST_CONTOH, data={"nama": "Token Mati", "_csrf": lama})
        self.assertEqual(j.status_code, 400)

    def test_tanpa_sesi_tetap_dialihkan_bukan_400(self):
        """Penjagaan masuk berjalan lebih dulu. Sesi yang habis di tengah
        pengisian formulir harus berakhir di halaman masuk, bukan di halaman
        galat yang buntu."""
        j = self.klien().post("/berkas/baru", data={"status": "draf", "_csrf": "ngawur"})
        self.assertEqual(j.status_code, 303)
        self.assertEqual(j.headers["Location"], "/masuk")

    def test_halaman_galat_menawarkan_jalan_keluar(self):
        isi = self.klien(masuk=True).post(self.POST_CONTOH,
                                          data={"_csrf": ""}).get_data(as_text=True)
        self.assertIn("kedaluwarsa", isi.lower())
        self.assertIn("Muat ulang", isi)

    # ----- penjaga cakupan
    def _forms_tanpa_token(self, isi):
        """Form POST yang tidak diikuti medan token tepat sesudah tag bukanya."""
        kurang = []
        for m in re.finditer(r'<form\b[^>]*\bmethod="post"[^>]*>', isi, re.I):
            ekor = isi[m.end():m.end() + 120]
            if 'name="_csrf"' not in ekor:
                kurang.append(m.group(0)[:70])
        return kurang

    def test_tiap_form_di_tiap_halaman_bertoken(self):
        c = self.klien(masuk=True)
        halaman = ["/masuk", "/berkas", "/berkas/baru", "/pradaftar", "/pradaftar/baru",
                   "/referensi", "/template", "/pengaturan"]
        halaman += [f"/referensi/bagian/{k}" for k in
                    ("jenis_hak", "klausa", "desa", "kecamatan", "kemendagri")]
        halaman += [f"/template/bagian/{k}" for k in ("cara", "arti", "penanda")]
        halaman += [f"/template/bagian/tpl-{j}" for j in
                    ("bap", "risalah", "sk", "bap-wakaf", "risalah-wakaf", "sk-wakaf")]
        n = 0
        for jalur in halaman:
            isi = c.get(jalur).get_data(as_text=True)
            n += len(re.findall(r'<form\b[^>]*\bmethod="post"', isi, re.I))
            self.assertEqual(self._forms_tanpa_token(isi), [], jalur)
        self.assertGreater(n, 15, "terlalu sedikit form terperiksa; daftarnya kadaluwarsa?")

    def test_form_pada_halaman_petugas_juga_bertoken(self):
        a = self.klien(masuk=True)
        a.post("/pengguna", data={"nama": "Csrf Uji", "username": "csrfuji",
                                  "sandi": "rahasia123", "peran": "petugas"})
        p = self.app.test_client()
        p.get("/masuk")
        p.post("/masuk", data={"username": "csrfuji", "sandi": "rahasia123"})
        for jalur in ("/pengaturan", "/berkas/baru", "/pradaftar/baru", "/referensi"):
            self.assertEqual(self._forms_tanpa_token(p.get(jalur).get_data(as_text=True)),
                             [], jalur)

    def test_form_pada_berkas_yang_sudah_ada_bertoken(self):
        c = self.klien(masuk=True)
        j = c.post("/berkas/baru", data={"status": "draf", "jenis_kegiatan": "baru"})
        bid = j.headers["Location"].split("/berkas/")[1].split("?")[0]
        isi = c.get(f"/berkas/{bid}").get_data(as_text=True)
        self.assertEqual(self._forms_tanpa_token(isi), [])
        self.assertIn('name="_csrf"', isi.split('id="form-berkas"')[1][:200])


class BerkasBaru(Dasar):
    """Formulir berkas baru dituntun per langkah; berkas yang sudah ada tidak."""

    def _panel(self, isi):
        return re.findall(r'<div class="panel" id="([^"]+)"', isi)

    def _tombol(self, isi):
        return re.findall(r'role="tab"[^>]*data-panel="([^"]+)"', isi)

    def test_enam_langkah_berurutan(self):
        isi = self.klien(masuk=True).get("/berkas/baru").get_data(as_text=True)
        urut = ["p-berkas", "p-pihak", "p-tanah", "p-asal", "p-dokumen", "p-foto"]
        self.assertEqual(self._tombol(isi), urut)
        self.assertEqual(self._panel(isi), urut)
        self.assertIn('class="tab tahap"', isi)

    def test_sidang_dan_cetak_tidak_ikut_saat_membuat(self):
        """Keduanya baru berarti setelah berkasnya ada; memunculkannya di sini
        hanya menambah langkah yang tak bisa diisi."""
        isi = self.klien(masuk=True).get("/berkas/baru").get_data(as_text=True)
        self.assertNotIn('id="p-sidang"', isi)
        self.assertNotIn('id="p-terbit"', isi)

    def test_berkas_yang_sudah_ada_tetap_delapan_tab(self):
        c = self.klien(masuk=True)
        j = c.post("/berkas/baru", data={"status": "draf", "penerima_nama": "Tab Lengkap"})
        bid = j.headers["Location"].split("/berkas/")[1].split("?")[0]
        isi = c.get(f"/berkas/{bid}").get_data(as_text=True)
        self.assertNotIn('class="tab tahap"', isi)
        self.assertEqual(len(self._panel(isi)), 8)
        self.assertIn('id="p-sidang"', isi)
        self.assertIn('id="p-terbit"', isi)

    def test_tiap_langkah_punya_petunjuk_dan_tombol_maju(self):
        isi = self.klien(masuk=True).get("/berkas/baru").get_data(as_text=True)
        self.assertEqual(isi.count("tahap-petunjuk"), 6)
        # lima langkah pertama punya Lanjut + Lewati, yang terakhir punya Simpan
        self.assertEqual(isi.count("data-tahap-maju"), 10)
        self.assertEqual(isi.count("data-tahap-mundur"), 5)
        self.assertEqual(isi.count("Simpan berkas"), 1)

    def test_tiga_langkah_akhir_ditandai_boleh_menyusul(self):
        isi = self.klien(masuk=True).get("/berkas/baru").get_data(as_text=True)
        bilah = isi.split('class="tab tahap"')[1].split("</div>")[0]
        self.assertEqual(bilah.count("boleh nanti"), 3)

    def test_tanpa_tombol_simpan_di_kepala_saat_membuat(self):
        """Satu-satunya Simpan ada di ujung langkah terakhir. Kalau ada juga di
        kepala, petugas baru cenderung menekannya di langkah pertama."""
        isi = self.klien(masuk=True).get("/berkas/baru").get_data(as_text=True)
        kepala = isi.split('class="lengket-atas"')[1].split("</div>")[2]
        self.assertNotIn("Simpan", kepala)

    def test_ingatan_langkah_terpisah_dari_ingatan_tab(self):
        """Kalau kuncinya sama, membuka berkas baru bisa mendarat di tengah
        wizard karena tab terakhir pada berkas lain."""
        isi = self.klien(masuk=True).get("/berkas/baru").get_data(as_text=True)
        self.assertIn('data-tab="berkas-baru"', isi)
        self.assertIn('data-tahap="1"', isi)

    def test_isian_wizard_benar_benar_tersimpan(self):
        c = self.klien(masuk=True)
        j = c.post("/berkas/baru", data={
            "jenis_hak_dimohon": "HM", "jenis_kegiatan": "baru", "status": "draf",
            "nomor_berkas": "9100/2026", "penerima_nama": "Siti Aminah",
            "penerima_nik": "7503016001900001"})
        self.assertEqual(j.status_code, 303)
        bid = j.headers["Location"].split("/berkas/")[1].split("?")[0]
        self.assertIn("pesan=dibuat", j.headers["Location"])
        isi = c.get(f"/berkas/{bid}").get_data(as_text=True)
        self.assertIn("Siti Aminah", isi)
        self.assertIn("9100/2026", isi)

    def test_langkah_yang_tak_diisi_tidak_bikin_gagal(self):
        """Tiga langkah terakhir boleh dilewati; simpan harus tetap jalan."""
        j = self.klien(masuk=True).post("/berkas/baru",
                                        data={"penerima_nama": "Lewat Saja"})
        self.assertEqual(j.status_code, 303)
        self.assertIn("/berkas/", j.headers["Location"])

    def test_petugas_juga_dapat_tuntunan(self):
        a = self.klien(masuk=True)
        a.post("/pengguna", data={"nama": "Baru Uji", "username": "baruuji",
                                  "sandi": "rahasia123", "peran": "petugas"})
        p = self.app.test_client()
        p.get("/masuk")
        p.post("/masuk", data={"username": "baruuji", "sandi": "rahasia123"})
        isi = p.get("/berkas/baru").get_data(as_text=True)
        self.assertIn('class="tab tahap"', isi)
        self.assertEqual(len(self._tombol(isi)), 6)

    def test_berkas_statis_bercap_versi(self):
        """Tanpa cap, peramban bisa memakai app.js lama bersama halaman baru —
        dan halaman setengah rusak itu sulit dikenali sebagai masalah singgahan."""
        isi = self.klien(masuk=True).get("/berkas/baru").get_data(as_text=True)
        self.assertRegex(isi, r'/static/style\.css\?v=\d+')
        self.assertRegex(isi, r'/static/app\.js\?v=\d+')


class Statis(Dasar):
    def test_dilayani_flask_tanpa_perlu_masuk(self):
        c = self.klien()
        for nama, tipe in (("style.css", "text/css"), ("app.js", "javascript")):
            j = c.get(f"/static/{nama}")
            self.assertEqual(j.status_code, 200, nama)
            self.assertIn(tipe, j.headers["Content-Type"], nama)
            j.close()

    def test_berkas_tak_ada_404(self):
        self.assertEqual(self.klien().get("/static/tidak-ada.css").status_code, 404)


class Referensi(Dasar):
    """Blueprint rute/referensi.py."""

    def _id_kecamatan(self, c, nama):
        """Ambil id kecamatan lewat tabelnya — sengaja bukan lewat SQL, supaya
        yang diuji ikut jalur yang benar-benar dipakai peramban."""
        isi = c.get("/referensi/bagian/kecamatan").get_data(as_text=True)
        for potong in isi.split("<form")[1:]:
            if nama in potong:
                m = re.search(r'name="id" value="(\d+)"', potong)
                if m:
                    return m.group(1)
        self.fail(f"kecamatan {nama!r} tidak ketemu di tabel")
        return None

    def test_halaman_terbuka(self):
        j = self.klien(masuk=True).get("/referensi")
        self.assertEqual(j.status_code, 200)
        self.assertIn("<!doctype html", j.get_data(as_text=True).lower())

    def test_tiap_kartu_lipat_terakit(self):
        c = self.klien(masuk=True)
        for kunci in ("jenis_hak", "klausa", "desa", "kecamatan", "kemendagri"):
            j = c.get(f"/referensi/bagian/{kunci}")
            self.assertEqual(j.status_code, 200, kunci)
            self.assertEqual(j.headers.get("X-Potongan"), "1", kunci)
            self.assertTrue(j.get_data(as_text=True).strip(), kunci)

    def test_kartu_tak_dikenal_404(self):
        self.assertEqual(
            self.klien(masuk=True).get("/referensi/bagian/entah").status_code, 404)

    def test_buka_kartu_lewat_query_tanpa_javascript(self):
        # ?buka= merakit isi kartunya di server; ini jalur cadangan kalau JS mati.
        j = self.klien(masuk=True).get("/referensi?buka=kecamatan")
        self.assertEqual(j.status_code, 200)
        self.assertIn("b-kecamatan", j.get_data(as_text=True))

    # ----- kecamatan
    def test_tambah_ubah_hapus_kecamatan(self):
        c = self.klien(masuk=True)
        j = c.post("/referensi/kecamatan", data={"nama": "Uji Kecamatan", "kode": "99.99.99"})
        self.assertEqual(j.status_code, 303)
        self.assertIn("pesan=kecamatan", j.headers["Location"])
        kid = self._id_kecamatan(c, "Uji Kecamatan")

        c.post("/referensi/kecamatan", data={"id": kid, "nama": "Uji Kecamatan Baru"})
        self.assertIn("Uji Kecamatan Baru",
                      c.get("/referensi/bagian/kecamatan").get_data(as_text=True))

        j = c.post("/referensi/kecamatan/hapus", data={"id": kid})
        self.assertIn("pesan=kecamatan_hapus", j.headers["Location"])
        self.assertNotIn("Uji Kecamatan Baru",
                         c.get("/referensi/bagian/kecamatan").get_data(as_text=True))

    def test_nama_kecamatan_wajib(self):
        j = self.klien(masuk=True).post("/referensi/kecamatan", data={"nama": "  "})
        self.assertIn("galat=", j.headers["Location"])

    def test_kecamatan_kembar_ditolak(self):
        c = self.klien(masuk=True)
        c.post("/referensi/kecamatan", data={"nama": "Kembar"})
        j = c.post("/referensi/kecamatan", data={"nama": "Kembar"})
        self.assertIn("pesan=kecamatan_ada", j.headers["Location"])

    def test_kecamatan_berisi_desa_tidak_bisa_dihapus(self):
        c = self.klien(masuk=True)
        c.post("/referensi/kecamatan", data={"nama": "Berisi"})
        kid = self._id_kecamatan(c, "Berisi")
        c.post("/referensi/desa", data={"nama": "Desa Uji", "kecamatan_id": kid})
        j = c.post("/referensi/kecamatan/hapus", data={"id": kid})
        self.assertIn("pesan=kecamatan_isi", j.headers["Location"])

    # ----- desa
    def test_tambah_desa_lalu_buka_halamannya(self):
        c = self.klien(masuk=True)
        c.post("/referensi/kecamatan", data={"nama": "Kec Desa"})
        kid = self._id_kecamatan(c, "Kec Desa")
        j = c.post("/referensi/desa", data={"nama": "Wontogia Uji", "kecamatan_id": kid,
                                            "jenis": "Desa"})
        self.assertEqual(j.status_code, 303)
        self.assertIn("/referensi/desa/", j.headers["Location"])
        j = c.get(j.headers["Location"].split("?")[0])
        self.assertEqual(j.status_code, 200)
        self.assertIn("Wontogia Uji", j.get_data(as_text=True))

    def test_desa_tanpa_kecamatan_ditolak(self):
        j = self.klien(masuk=True).post("/referensi/desa", data={"nama": "Tanpa Kecamatan"})
        self.assertIn("galat=", j.headers["Location"])

    def test_desa_kembar_di_kecamatan_yang_sama_ditolak(self):
        c = self.klien(masuk=True)
        c.post("/referensi/kecamatan", data={"nama": "Kec Kembar"})
        kid = self._id_kecamatan(c, "Kec Kembar")
        c.post("/referensi/desa", data={"nama": "Sama", "kecamatan_id": kid})
        j = c.post("/referensi/desa", data={"nama": "Sama", "kecamatan_id": kid})
        self.assertIn("pesan=desa_ada", j.headers["Location"])

    def test_desa_tak_ada_dialihkan_ke_daftar(self):
        j = self.klien(masuk=True).get("/referensi/desa/999999")
        self.assertEqual(j.status_code, 303)
        self.assertIn("buka=desa", j.headers["Location"])

    def test_kepala_desa_disimpan(self):
        c = self.klien(masuk=True)
        c.post("/referensi/kecamatan", data={"nama": "Kec Pejabat"})
        kid = self._id_kecamatan(c, "Kec Pejabat")
        j = c.post("/referensi/desa", data={"nama": "Desa Pejabat", "kecamatan_id": kid})
        did = j.headers["Location"].split("/referensi/desa/")[1].split("?")[0]

        j = c.post("/referensi/desa/pejabat",
                   data={"desa_id": did, "nama": "Hasan Ali", "aktif": "1"})
        self.assertIn("pesan=pejabat", j.headers["Location"])
        self.assertIn("Hasan Ali", c.get(f"/referensi/desa/{did}").get_data(as_text=True))

    def test_nama_kepala_desa_wajib(self):
        j = self.klien(masuk=True).post("/referensi/desa/pejabat", data={"desa_id": "1"})
        self.assertIn("galat=", j.headers["Location"])

    # ----- susunan panitia
    def test_daur_hidup_susunan_panitia(self):
        c = self.klien(masuk=True)
        j = c.post("/referensi/panitia", data={"nomor_sk": "99/SK-UJI/2026",
                                               "tanggal_sk": "2026-01-05"})
        self.assertEqual(j.status_code, 303)
        self.assertIn("buka=panitia-", j.headers["Location"])
        pid = j.headers["Location"].split("buka=panitia-")[1].split("#")[0]

        j = c.post("/referensi/panitia/anggota",
                   data={"panitia_id": pid, "nama": "Sitti Rahma", "jabatan": "Ketua",
                         "urut": "1"})
        self.assertIn("pesan=anggota", j.headers["Location"])
        self.assertIn("Sitti Rahma",
                      c.get(f"/referensi/bagian/panitia-{pid}").get_data(as_text=True))

        j = c.post("/referensi/panitia/salin", data={"panitia_id": pid})
        self.assertIn("pesan=panitia_salin", j.headers["Location"])
        salinan = j.headers["Location"].split("buka=panitia-")[1].split("#")[0]
        self.assertNotEqual(salinan, pid, "salinan harus jadi baris baru")
        self.assertIn("Sitti Rahma",
                      c.get(f"/referensi/bagian/panitia-{salinan}").get_data(as_text=True))

        j = c.post("/referensi/panitia/hapus", data={"panitia_id": salinan})
        self.assertIn("pesan=panitia_hapus", j.headers["Location"])
        self.assertEqual(c.get(f"/referensi/bagian/panitia-{salinan}").status_code, 404)

    def test_nomor_sk_panitia_wajib(self):
        j = self.klien(masuk=True).post("/referensi/panitia", data={"nomor_sk": ""})
        self.assertIn("galat=", j.headers["Location"])

    def test_nama_anggota_wajib(self):
        j = self.klien(masuk=True).post("/referensi/panitia/anggota",
                                        data={"panitia_id": "1", "nama": ""})
        self.assertIn("galat=", j.headers["Location"])

    def test_selaraskan_daftar_kemendagri(self):
        j = self.klien(masuk=True).post("/referensi/wilayah")
        self.assertEqual(j.status_code, 303)
        self.assertTrue(j.headers["Location"].startswith("/referensi"))


class Template(Dasar):
    """Blueprint rute/template.py.

    Yang diuji hanya jalur baca dan jalur tolak: template sungguhan di
    templates/ tidak boleh sampai tertimpa oleh uji.
    """

    def _cap_template(self):
        p = templat.jalur("sk")
        return os.path.getsize(p), os.path.getmtime(p)

    def test_halaman_terbuka(self):
        j = self.klien(masuk=True).get("/template")
        self.assertEqual(j.status_code, 200)
        self.assertIn("<!doctype html", j.get_data(as_text=True).lower())

    def test_kartu_cara_dan_arti_terakit(self):
        c = self.klien(masuk=True)
        for kunci in ("cara", "arti"):
            j = c.get(f"/template/bagian/{kunci}")
            self.assertEqual(j.status_code, 200, kunci)
            self.assertEqual(j.headers.get("X-Potongan"), "1", kunci)

    def test_kartu_tak_dikenal_404(self):
        self.assertEqual(
            self.klien(masuk=True).get("/template/bagian/entah").status_code, 404)

    def test_unduh_template_docx(self):
        j = self.klien(masuk=True).get("/template/unduh/sk")
        self.assertEqual(j.status_code, 200)
        self.assertIn("wordprocessingml", j.headers["Content-Type"])
        self.assertIn("template-sk.docx", j.headers["Content-Disposition"])
        self.assertEqual(j.get_data()[:2], b"PK")        # docx = arsip zip
        j.close()

    def test_unduh_jenis_tak_dikenal_404(self):
        self.assertEqual(
            self.klien(masuk=True).get("/template/unduh/ngawur").status_code, 404)

    def test_unggah_tanpa_berkas_ditolak(self):
        j = self.klien(masuk=True).post("/template/unggah", data={"jenis": "sk"})
        self.assertIn("galat=", j.headers["Location"])

    def test_unggah_jenis_tak_dikenal_ditolak(self):
        j = self.klien(masuk=True).post("/template/unggah", data={"jenis": "ngawur"})
        self.assertIn("galat=", j.headers["Location"])

    def test_unggah_bukan_docx_ditolak_sebelum_menimpa(self):
        c = self.klien(masuk=True)
        sebelum = self._cap_template()
        j = c.post("/template/unggah",
                   data={"jenis": "sk", "berkas": (io.BytesIO(b"bukan docx"), "catatan.txt")},
                   content_type="multipart/form-data")
        self.assertIn("galat=", j.headers["Location"])
        self.assertEqual(self._cap_template(), sebelum,
                         "template asli ikut tertimpa oleh unggahan yang ditolak")


if __name__ == "__main__":
    unittest.main(verbosity=2)


class Pradaftar(Dasar):
    """Blueprint rute/pradaftar.py — loket, daftar kelengkapan, dan naik jadi berkas."""

    def _butir(self, c, pid):
        """Peta id butir yang bisa dicentang, dibaca dari halaman yang baru dirakit.

        Sengaja dari HTML, bukan dari basis data: yang diuji justru apakah
        butirnya benar-benar sampai ke formulir sebagai isian yang bisa dikirim.
        """
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        return [int(x) for x in re.findall(r'name="jawab_butir" value="(\d+)"', isi)]

    def _buat(self, c, **tambahan):
        data = dict(self.ISIAN)
        data.update(tambahan)
        j = c.post("/pradaftar/baru", data=data)
        self.assertEqual(j.status_code, 303)
        return int(j.headers["Location"].split("/pradaftar/")[1].split("?")[0])

    def _centang(self, c, pid, ids, **tambahan):
        data = dict(self.ISIAN, jawab_butir=[str(x) for x in ids])
        for x in ids:
            data[f"jawab_ada_{x}"] = "1"
        data.update(tambahan)
        return c.post(f"/pradaftar/{pid}", data=data)

    def _ditolak(self, j, pesan=""):
        """POST yang ditolak izin dialihkan, bukan dijawab 403 telanjang —
        lihat errorhandler(403) di aplikasi.py. Yang diperiksa bentuk itu."""
        self.assertEqual(j.status_code, 303, pesan)
        self.assertIn("galat=", j.headers["Location"], pesan)

    ISIAN = {"jenis_hak": "HM", "jenis_kegiatan": "baru",
             "asal_tanah": "tanah_negara", "pemohon_nama": "Usman Isa",
             "pemohon_jenis_subjek": "perorangan"}

    # ------------------------------------------------------------- halaman
    def test_daftar_terakit_penuh(self):
        j = self.klien(masuk=True).get("/pradaftar")
        self.assertEqual(j.status_code, 200)
        isi = j.get_data(as_text=True)
        self.assertIn("<h1>Pradaftar</h1>", isi)
        # tombol saringnya ikut, karena app.js menggantungkan penyaringan padanya
        self.assertIn('data-status="kurang"', isi)

    def test_pradaftar_baru_dituntun_berlangkah(self):
        isi = self.klien(masuk=True).get("/pradaftar/baru").get_data(as_text=True)
        for panel in ("q-pemohon", "q-tanah", "q-periksa"):
            self.assertIn(f'id="{panel}"', isi)
        # Tab Kesimpulan sengaja tidak ikut saat membuat: belum ada yang bisa
        # diterima, dan tombolnya akan menunjuk pradaftar yang belum ada.
        self.assertNotIn('id="q-putusan"', isi)

    def test_sesudah_disimpan_jadi_empat_tab(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn('data-tab="pradaftar"', isi)
        for panel in ("q-pemohon", "q-tanah", "q-periksa", "q-putusan"):
            self.assertIn(f'id="{panel}"', isi)

    def test_nomor_agenda_diberi_saat_dibuat_dan_berurutan(self):
        c = self.klien(masuk=True)
        a, b = self._buat(c), self._buat(c)
        cari = lambda pid: re.search(
            r"PD-(\d{4})/(\d{4})", c.get(f"/pradaftar/{pid}").get_data(as_text=True))
        ma, mb = cari(a), cari(b)
        self.assertIsNotNone(ma)
        self.assertEqual(int(mb.group(1)), int(ma.group(1)) + 1)
        self.assertEqual(ma.group(2), mb.group(2))

    # --------------------------------------------------- daftar kelengkapan
    def test_butir_hak_milik_terpasang_apa_adanya(self):
        c = self.klien(masuk=True)
        isi = c.get(f"/pradaftar/{self._buat(c)}").get_data(as_text=True)
        self.assertIn("Daftar Kelengkapan Persyaratan Permohonan Hak Milik", isi)
        for bunyi in ("Permohonan Hak Milik", "Identitas Pemohon",
                      "Dasar Penguasaan atau Alas Hak", "Peta Bidang Tanah",
                      "Bukti perpajakan", "Risalah lelang",
                      "Surat Pengantar dari Kantor Pertanahan"):
            self.assertIn(bunyi, isi, bunyi)

    def test_grup_badan_hukum_hanya_berlaku_untuk_badan_hukum(self):
        c = self.klien(masuk=True)
        orang = self._buat(c, pemohon_jenis_subjek="perorangan")
        badan = self._buat(c, pemohon_jenis_subjek="badan_hukum")
        self.assertLess(len(self._butir(c, orang)), len(self._butir(c, badan)))
        # Bunyinya tetap tercetak di kedua-duanya: formulir resminya memuatnya,
        # hanya kolom Ada/Tidak Ada-nya yang dikosongkan.
        self.assertIn("izin perolehan tanah",
                      c.get(f"/pradaftar/{orang}").get_data(as_text=True))

    def test_butir_hpl_berlaku_hanya_di_atas_hak_pengelolaan(self):
        c = self.klien(masuk=True)
        negara = self._buat(c, asal_tanah="tanah_negara")
        hpl = self._buat(c, asal_tanah="hak_pengelolaan")
        self.assertNotEqual(len(self._butir(c, negara)), len(self._butir(c, hpl)))

    def test_centangan_tersimpan_dan_terbaca_lagi(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        ids = self._butir(c, pid)
        self._centang(c, pid, ids[:1])
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn(f'name="jawab_ada_{ids[0]}" value="1"', isi)
        self.assertIn("checked", isi)

    def test_kekurangan_dihitung_ulang_bukan_disimpan(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        ids = self._butir(c, pid)
        self.assertIn("belum lengkap", c.get(f"/pradaftar/{pid}").get_data(as_text=True))
        self._centang(c, pid, ids)
        self.assertIn("Kelengkapan terpenuhi",
                      c.get(f"/pradaftar/{pid}").get_data(as_text=True))
        # dikosongkan lagi: kembali kurang, tanpa satu pun kolom status diubah
        self._centang(c, pid, [])
        self.assertIn("belum lengkap", c.get(f"/pradaftar/{pid}").get_data(as_text=True))

    def test_butir_isian_bebas_menyimpan_nama_suratnya(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        j_id = int(re.search(r'name="jawab_uraian_(\d+)"', isi).group(1))
        self._centang(c, pid, [j_id],
                      **{f"jawab_uraian_{j_id}": "Surat Keterangan Hibah 12/2019"})
        self.assertIn("Surat Keterangan Hibah 12/2019",
                      c.get(f"/pradaftar/{pid}").get_data(as_text=True))

    # ------------------------------------------------ periksa layar penuh
    def test_belum_diperiksa_tidak_otomatis_tidak_ada(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        ids = self._butir(c, pid)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        # Pradaftar baru: tak satu pun pilihan tercentang, semuanya "belum".
        self.assertNotIn(" checked", isi.split('class="periksa"')[1])
        self.assertIn(f"0</b> dari", isi)
        # "Tidak Ada" yang dipilih sendiri tersimpan sebagai jawaban.
        self._centang(c, pid, [ids[0]], **{f"jawab_ada_{ids[0]}": "0"})
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertRegex(isi, rf'name="jawab_ada_{ids[0]}" value="0" checked')
        self.assertIn("1</b> dari", isi)

    def test_layar_penuh_tampil_tanpa_menu_samping(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.get(f"/pradaftar/{pid}/periksa")
        self.assertEqual(j.status_code, 200)
        isi = j.get_data(as_text=True)
        self.assertIn('<main class="fokus">', isi)
        self.assertIn(f'action="/pradaftar/{pid}/periksa"', isi)
        self.assertEqual(len(re.findall(r'name="jawab_butir"', isi)),
                         len(self._butir(c, pid)))

    def test_simpan_layar_penuh_tidak_menimpa_data_pemohon(self):
        c = self.klien(masuk=True)
        pid = self._buat(c, pemohon_nik="7503010101900001")
        ids = self._butir(c, pid)
        data = {"jawab_butir": [str(x) for x in ids]}
        data.update({f"jawab_ada_{x}": "1" for x in ids})
        j = c.post(f"/pradaftar/{pid}/periksa", data=data)
        self.assertEqual(j.status_code, 303)
        self.assertIn(f"/pradaftar/{pid}/periksa", j.headers["Location"])
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn("Usman Isa", isi)
        self.assertIn("7503010101900001", isi)
        self.assertIn("Kelengkapan terpenuhi", isi)

    def test_simpan_dan_kembali_dari_layar_penuh(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.post(f"/pradaftar/{pid}/periksa", data={"lanjut": "kembali"})
        self.assertTrue(j.headers["Location"].endswith(f"/pradaftar/{pid}?pesan=tersimpan#q-periksa"))

    def test_tombol_layar_penuh_menyimpan_dulu(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = self._centang(c, pid, [], lanjut="periksa")
        self.assertTrue(j.headers["Location"].endswith(f"/pradaftar/{pid}/periksa"))

    def test_layar_penuh_yang_terkunci_baca_saja(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        c.post(f"/pradaftar/{pid}/status", data={"status": "batal"})
        isi = c.get(f"/pradaftar/{pid}/periksa").get_data(as_text=True)
        self.assertIn("hanya bisa dilihat", isi)
        self.assertNotIn("data-simpan-pintas", isi)
        self._ditolak(c.post(f"/pradaftar/{pid}/periksa", data={}))

    # -------------------------------------------------------- terima loket
    def test_yang_masih_kurang_tidak_bisa_diterima_tanpa_alasan(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.post(f"/pradaftar/{pid}/terima", data={})
        self.assertEqual(j.status_code, 303)
        self.assertIn(f"/pradaftar/{pid}", j.headers["Location"])
        self.assertIn("galat=", j.headers["Location"])

    def test_terima_bersyarat_dengan_alasan_boleh(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.post(f"/pradaftar/{pid}/terima",
                   data={"alasan_terima": "kekurangan menyusul 3 hari kerja"})
        self.assertIn("/berkas/", j.headers["Location"])

    def test_yang_lengkap_naik_jadi_berkas_berisi_pemohon_dan_dokumen(self):
        c = self.klien(masuk=True)
        tambah = {"pemohon_nik": "7503010101900001", "pemohon_alamat": "Desa Tupa",
                  "nomor_pbt": "00123/2026", "luas_surat": "500"}
        pid = self._buat(c, **tambah)
        self._centang(c, pid, self._butir(c, pid), **tambah)
        j = c.post(f"/pradaftar/{pid}/terima", data={})
        self.assertEqual(j.status_code, 303)
        bid = j.headers["Location"].split("/berkas/")[1].split("?")[0]
        isi = c.get(f"/berkas/{bid}").get_data(as_text=True)
        self.assertIn("Usman Isa", isi)
        self.assertIn("7503010101900001", isi)
        self.assertIn("00123/2026", isi)
        # nama surat yang tadi dicentang ikut terbawa ke slot dokumen bakunya
        self.assertIn("Surat permohonan Hak Milik", isi)

    def test_tidak_bisa_diterima_dua_kali(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        self._centang(c, pid, self._butir(c, pid))
        pertama = c.post(f"/pradaftar/{pid}/terima").headers["Location"]
        self.assertIn("/berkas/", pertama)
        # Sekali terkunci, rute menolaknya sebelum menyentuh basis data —
        # nomor berkas yang telanjur keluar tidak bisa ditarik lagi, dan berkas
        # kembar berarti satu permohonan disidangkan dua kali.
        self._ditolak(c.post(f"/pradaftar/{pid}/terima"))
        self.assertEqual(c.get("/pradaftar").get_data(as_text=True).count(
            f'href="/berkas/{pertama.split("/berkas/")[1].split("?")[0]}"'), 1)

    def test_yang_sudah_diterima_jadi_baca_saja(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        self._centang(c, pid, self._butir(c, pid))
        c.post(f"/pradaftar/{pid}/terima")
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn("Sudah diterima di loket", isi)
        self.assertIn("<fieldset disabled", isi)
        self._ditolak(self._centang(c, pid, []))

    def test_status_dikembalikan_masih_bisa_disunting_batal_tidak(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.post(f"/pradaftar/{pid}/status", data={"status": "dikembalikan"})
        self.assertEqual(j.status_code, 303)
        self.assertIn("Dikembalikan", c.get("/pradaftar").get_data(as_text=True))
        # Pemohon datang lagi melengkapi, jadi yang dikembalikan tetap terbuka.
        self.assertEqual(self._centang(c, pid, []).status_code, 303)
        c.post(f"/pradaftar/{pid}/status", data={"status": "batal"})
        self._ditolak(self._centang(c, pid, []))

    def test_status_yang_tak_dikenal_ditolak(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.post(f"/pradaftar/{pid}/status", data={"status": "ngawur"})
        self.assertEqual(j.status_code, 400)


    # -------------------------------------------------------------- cetak
    def _docx(self, j):
        """Teks seluruh paragraf dan sel tabel dari jawaban DOCX."""
        from docx import Document
        if j.status_code != 200:              # badannya teks galat, bukan DOCX
            self.fail(f"{j.status_code}: {j.get_data(as_text=True)[:200]}")
        self.assertIn("wordprocessingml", j.headers["Content-Type"])
        d = Document(io.BytesIO(j.get_data()))
        teks = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                teks += [c.text for c in r.cells]
        return "\n".join(teks)

    def test_cetak_daftar_kelengkapan_memuat_seluruh_butir(self):
        c = self.klien(masuk=True)
        pid = self._buat(c, pemohon_nik="7503010101900001", pemohon_alamat="Desa Tupa")
        self._centang(c, pid, self._butir(c, pid)[:3],
                      pemohon_nik="7503010101900001", pemohon_alamat="Desa Tupa")
        isi = self._docx(c.get(f"/pradaftar/{pid}/docx/kelengkapan"))
        self.assertIn("Daftar Kelengkapan Persyaratan Permohonan Hak Milik", isi)
        self.assertIn("Usman Isa", isi)
        self.assertIn("7503010101900001", isi)
        # butir yang tidak berlaku pun tetap tercetak — formulir resmi memuatnya
        self.assertIn("izin perolehan tanah", isi)
        self.assertIn("Surat Pengantar dari Kantor Pertanahan", isi)
        self.assertIn("Coret yang tidak perlu", isi)
        # tak boleh ada penanda yang lolos tanpa diganti
        self.assertNotIn("{{", isi)

    def test_cetak_memberi_nama_berkas_bernomor_pradaftar(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.get(f"/pradaftar/{pid}/docx/kelengkapan")
        self.assertIn("Kelengkapan", j.headers["Content-Disposition"])
        self.assertIn("PD-", j.headers["Content-Disposition"])

    def test_surat_pengembalian_hanya_memuat_yang_kurang(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        isi = self._docx(c.get(f"/pradaftar/{pid}/docx/pengembalian"))
        self.assertIn("PENGEMBALIAN BERKAS PERMOHONAN", isi)
        self.assertIn("Dasar Penguasaan atau Alas Hak", isi)
        # yang bukan kekurangan tidak ikut: surat ini daftar kekurangan, bukan
        # salinan seluruh formulir
        self.assertNotIn("Risalah lelang", isi)
        self.assertNotIn("{{", isi)

    def test_surat_pengembalian_ditolak_bila_sudah_lengkap(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        self._centang(c, pid, self._butir(c, pid))
        j = c.get(f"/pradaftar/{pid}/docx/pengembalian")
        self.assertEqual(j.status_code, 400)
        self.assertIn("sudah terpenuhi", j.get_data(as_text=True))

    def test_jenis_cetakan_tak_dikenal_ditolak(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        for jalur in ("docx", "pdf", "pratinjau"):
            j = c.get(f"/pradaftar/{pid}/{jalur}/ngawur")
            self.assertEqual(j.status_code, 400, jalur)

    # ------------------------------------------------------------ pratinjau
    def test_tombol_cetak_menunjuk_ke_pratinjau_bukan_unduhan(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn(f'href="/pradaftar/{pid}/pratinjau/kelengkapan"', isi)
        self.assertIn(f'href="/pradaftar/{pid}/pratinjau/pengembalian"', isi)

    def test_pratinjau_menyediakan_cetak_pdf_dan_docx(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.get(f"/pradaftar/{pid}/pratinjau/kelengkapan")
        self.assertEqual(j.status_code, 200)
        isi = j.get_data(as_text=True)
        # Halamannya sama untuk komputer yang punya mesin PDF maupun tidak;
        # yang tanpa mesin menampilkan kabar galat plus tautan unduh DOCX.
        self.assertIn(f"/pradaftar/{pid}/docx/kelengkapan", isi)
        if "bingkai-pdf" in isi:
            self.assertIn(f'src="/pradaftar/{pid}/pdf/kelengkapan"', isi)
            self.assertIn('data-cetak="bingkai-pdf"', isi)
            self.assertIn(f"/pradaftar/{pid}/pdf/kelengkapan?unduh=1", isi)
            self.assertIn("segarkan=1", isi)
        else:
            self.assertIn("Microsoft Word atau LibreOffice", isi)

    def test_pratinjau_surat_pengembalian_yang_sudah_lengkap_ditolak_berhalaman(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        self._centang(c, pid, self._butir(c, pid))
        j = c.get(f"/pradaftar/{pid}/pratinjau/pengembalian")
        self.assertEqual(j.status_code, 400)
        self.assertIn("sudah terpenuhi", j.get_data(as_text=True))

    def test_pdf_dilewati_rapi_bila_tak_ada_mesin(self):
        """Di komputer tanpa Word maupun LibreOffice, PDF-nya dijawab 501 —
        bukan galat 500 yang tidak bisa dibaca petugas."""
        from berkas.dokumen import pdf as modul_pdf
        c = self.klien(masuk=True)
        pid = self._buat(c)
        asli = modul_pdf.mesin_tersedia
        modul_pdf.mesin_tersedia = lambda: None
        try:
            j = c.get(f"/pradaftar/{pid}/pdf/kelengkapan")
            self.assertEqual(j.status_code, 501)
            self.assertIn("LibreOffice", j.get_data(as_text=True))
            # halaman pratinjaunya tetap terakit, dengan jalan keluar DOCX
            isi = c.get(f"/pradaftar/{pid}/pratinjau/kelengkapan").get_data(as_text=True)
            self.assertIn(f"/pradaftar/{pid}/docx/kelengkapan", isi)
        finally:
            modul_pdf.mesin_tersedia = asli

    def test_singgahan_docx_ikut_waktu_ubah_pradaftar(self):
        """Umur singgahan ditentukan datanya, bukan waktu perakitan.

        Itulah yang membuat PDF-nya bisa dipakai ulang tanpa mencatat apa pun:
        pdf.ke_pdf() membandingkan mtime, dan mtime DOCX-nya sengaja disetel ke
        waktu ubah pradaftarnya. Merakit ulang dokumen yang sama tidak boleh
        menggeser mtime-nya — kalau bergeser, tiap pratinjau memanggil Word lagi.
        """
        import os
        from berkas import db as modul_db
        from berkas.dokumen import kelengkapan as modul, templat
        import urllib.parse as up
        c = self.klien(masuk=True)
        pid = self._buat(c)
        j = c.get(f"/pradaftar/{pid}/docx/kelengkapan")
        # Nama berkasnya diambil dari jawabannya sendiri: uji lain di kelas ini
        # ikut meninggalkan singgahan di folder yang sama.
        nama = up.unquote(j.headers["Content-Disposition"].split("UTF-8''")[1])
        jalur = os.path.join(modul.pdf.DIR_PRATINJAU, nama)
        self.assertTrue(os.path.isfile(jalur), f"{nama} tidak ada di pratinjau/")

        k = modul_db.sambung()
        r = k.execute("SELECT dibuat, diubah FROM pradaftar WHERE id=?", (pid,)).fetchone()
        k.close()
        harap = max(modul._waktu(r["diubah"]), modul._waktu(r["dibuat"]),
                    os.path.getmtime(templat.jalur("checklist")))
        self.assertEqual(os.path.getmtime(jalur), harap)

        c.get(f"/pradaftar/{pid}/docx/kelengkapan")       # dirakit ulang
        self.assertEqual(os.path.getmtime(jalur), harap, "mtime tidak boleh bergeser")

    def test_yang_sudah_diterima_tetap_boleh_dicetak(self):
        # Formulirnya justru dibutuhkan sebagai lampiran arsip berkas, jadi
        # penguncian sunting tidak ikut mengunci pencetakan.
        c = self.klien(masuk=True)
        pid = self._buat(c)
        self._centang(c, pid, self._butir(c, pid))
        c.post(f"/pradaftar/{pid}/terima")
        isi = self._docx(c.get(f"/pradaftar/{pid}/docx/kelengkapan"))
        self.assertIn("Usman Isa", isi)
        # tombolnya pun masih ada di halaman
        self.assertIn(f'href="/pradaftar/{pid}/pratinjau/kelengkapan"',
                      c.get(f"/pradaftar/{pid}").get_data(as_text=True))

    def test_cetakan_masuk_singgahan_bukan_folder_keluaran(self):
        """Cetakan loket bisa dirakit ulang persis dari pradaftarnya, jadi
        tempatnya di pratinjau/ yang dipangkas sendiri — bukan di keluaran/
        yang isinya naskah resmi, dan bukan pula di arsip cetak."""
        import os
        from berkas.dokumen import terbitkan
        c = self.klien(masuk=True)
        pid = self._buat(c)
        c.get(f"/pradaftar/{pid}/docx/kelengkapan")
        keluaran = (os.listdir(terbitkan.DIR_KELUARAN)
                    if os.path.isdir(terbitkan.DIR_KELUARAN) else [])
        self.assertEqual([x for x in keluaran if "Kelengkapan" in x], [])
        self.assertNotIn("/unduh/", c.get(f"/pradaftar/{pid}").get_data(as_text=True))

    def test_petugas_lain_tak_boleh_mencetak(self):
        a = self.klien(masuk=True)
        pid = self._buat(a)
        p = self._petugas(a, "loketcetak")
        for jalur in ("docx", "pdf", "pratinjau"):
            j = p.get(f"/pradaftar/{pid}/{jalur}/kelengkapan")
            self.assertEqual(j.status_code, 403, jalur)

    def test_berkas_menunjuk_balik_ke_pradaftar_asalnya(self):
        """Tautannya satu arah: berkas tidak menyimpan kolom pradaftar_id,
        yang tahu dirinya sudah naik adalah pradaftarnya."""
        c = self.klien(masuk=True)
        pid = self._buat(c)
        self._centang(c, pid, self._butir(c, pid))
        bid = c.post(f"/pradaftar/{pid}/terima").headers["Location"]                .split("/berkas/")[1].split("?")[0]
        isi = c.get(f"/berkas/{bid}").get_data(as_text=True)
        self.assertIn("berasal dari", isi)
        self.assertIn(f'href="/pradaftar/{pid}"', isi)

    def test_berkas_biasa_tidak_menyebut_pradaftar(self):
        c = self.klien(masuk=True)
        j = c.post("/berkas/baru", data={"status": "draf", "jenis_kegiatan": "baru"})
        bid = j.headers["Location"].split("/berkas/")[1].split("?")[0]
        self.assertNotIn("berasal dari", c.get(f"/berkas/{bid}").get_data(as_text=True))

    # --------------------------------------------------------------- izin
    def _petugas(self, admin, username):
        admin.post("/pengguna", data={"nama": "Loket " + username, "username": username,
                                      "sandi": "rahasia123", "peran": "petugas"})
        p = self.app.test_client()
        p.get("/masuk")
        p.post("/masuk", data={"username": username, "sandi": "rahasia123"})
        p.segarkan()
        return p

    def test_petugas_tak_melihat_pradaftar_orang_lain(self):
        a = self.klien(masuk=True)
        pid = self._buat(a)
        p = self._petugas(a, "loketuji")
        self.assertEqual(p.get(f"/pradaftar/{pid}").status_code, 403)
        self.assertNotIn("Usman Isa", p.get("/pradaftar").get_data(as_text=True))

    def test_petugas_tidak_boleh_menghapus(self):
        a = self.klien(masuk=True)
        pid = self._buat(a)
        p = self._petugas(a, "loketdua")
        j = p.post(f"/pradaftar/{pid}/hapus")
        self.assertNotIn("pradaftar_hapus", j.headers.get("Location", ""))
        self.assertEqual(a.get(f"/pradaftar/{pid}").status_code, 200)

    def test_admin_menghapus_pradaftar_tanpa_membawa_berkasnya(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        self._centang(c, pid, self._butir(c, pid))
        tujuan = c.post(f"/pradaftar/{pid}/terima").headers["Location"]
        bid = tujuan.split("/berkas/")[1].split("?")[0]
        j = c.post(f"/pradaftar/{pid}/hapus")
        self.assertEqual(j.headers["Location"], "/pradaftar?pesan=pradaftar_hapus")
        self.assertEqual(c.get(f"/pradaftar/{pid}").status_code, 404)
        self.assertEqual(c.get(f"/berkas/{bid}").status_code, 200)

    # ---------------------------------------------- daftar kelengkapan (referensi)
    def _formulir_id(self, c):
        isi = c.get("/referensi").get_data(as_text=True)
        return int(re.search(r'id="b-kelengkapan-(\d+)"', isi).group(1))

    def test_formulir_kelengkapan_tampil_di_data_referensi(self):
        c = self.klien(masuk=True)
        isi = c.get("/referensi").get_data(as_text=True)
        self.assertIn("Daftar Kelengkapan Persyaratan Permohonan Hak Milik", isi)
        fid = self._formulir_id(c)
        kartu = c.get(f"/referensi/bagian/kelengkapan-{fid}").get_data(as_text=True)
        self.assertIn("Dasar Penguasaan atau Alas Hak", kartu)
        self.assertIn("Surat Pengantar dari Kantor Pertanahan", kartu)

    def test_admin_boleh_mengubah_bunyi_butir_dan_wajibnya(self):
        c = self.klien(masuk=True)
        fid = self._formulir_id(c)
        kartu = c.get(f"/referensi/bagian/kelengkapan-{fid}").get_data(as_text=True)
        bid = int(re.search(r'name="nama_(\d+)"', kartu).group(1))
        j = c.post("/referensi/kelengkapan",
                   data={"id": str(fid), f"nama_{bid}": "Permohonan Hak Milik (asli)"})
        self.assertEqual(j.status_code, 303)
        kartu = c.get(f"/referensi/bagian/kelengkapan-{fid}").get_data(as_text=True)
        self.assertIn("Permohonan Hak Milik (asli)", kartu)
        # wajibnya ikut mati karena kotaknya tidak dikirim — itulah cara kerja
        # kotak centang, dan rute memang membacanya begitu
        self.assertNotIn(f'name="wajib_{bid}" value="1" checked', kartu)

    def test_bunyi_butir_tidak_boleh_dikosongkan(self):
        c = self.klien(masuk=True)
        fid = self._formulir_id(c)
        kartu = c.get(f"/referensi/bagian/kelengkapan-{fid}").get_data(as_text=True)
        bid = int(re.search(r'name="nama_(\d+)"', kartu).group(1))
        semula = re.search(r'name="nama_%d" value="([^"]*)"' % bid, kartu).group(1)
        c.post("/referensi/kelengkapan", data={"id": str(fid), f"nama_{bid}": "   "})
        kartu = c.get(f"/referensi/bagian/kelengkapan-{fid}").get_data(as_text=True)
        self.assertIn(semula, kartu)

    def test_petugas_tidak_boleh_mengubah_formulir(self):
        a = self.klien(masuk=True)
        fid = self._formulir_id(a)
        p = self._petugas(a, "loketref")
        j = p.post("/referensi/kelengkapan", data={"id": str(fid)})
        self._ditolak(j)

    # ------------------------------------------------------ perlu koreksi
    def _koreksi(self, c, pid, ids, bid, catatan):
        """Semua butir dicentang Ada, kecuali `bid` yang ada tapi perlu koreksi."""
        return self._centang(c, pid, ids, **{f"jawab_ada_{bid}": "3",
                                             f"jawab_catatan_{bid}": catatan})

    def test_perlu_koreksi_menahan_penerimaan_dan_catatannya_tampil(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        ids = self._butir(c, pid)
        self._koreksi(c, pid, ids, ids[0], "Belum dilegalisir camat")
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn("1 perlu koreksi", isi)
        self.assertIn("Belum dilegalisir camat", isi)
        self.assertNotIn("Kelengkapan terpenuhi", isi)
        self.assertIn(f'name="jawab_ada_{ids[0]}" value="3" checked', isi)
        # tanpa alasan tertulis ditolak, sama seperti yang kurang
        j = c.post(f"/pradaftar/{pid}/terima", data={})
        self.assertIn("galat=", j.headers["Location"])
        self.assertIn("koreksi", c.get("/pradaftar").get_data(as_text=True))

    def test_catatan_dibuang_bila_pilihannya_bukan_koreksi(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        ids = self._butir(c, pid)
        # kotak catatan tetap terkirim walau pilihannya sudah diganti ke Ada
        self._centang(c, pid, ids, **{f"jawab_catatan_{ids[0]}": "catatan basi"})
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertNotIn("catatan basi", isi)
        self.assertIn("Kelengkapan terpenuhi", isi)

    def test_koreksi_pada_alternatif_tidak_membuat_grupnya_kurang(self):
        """Alas hak yang dibawa tapi salah: yang diminta perbaikan, bukan 'bawa
        salah satu dari a sampai j'."""
        c = self.klien(masuk=True)
        pid = self._buat(c)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        ids = self._butir(c, pid)
        hibah = int(re.search(r'name="jawab_uraian_(\d+)"', isi).group(1))
        # semua Tidak Ada, kecuali satu alternatif alas hak yang perlu koreksi
        data = dict(self.ISIAN, jawab_butir=[str(x) for x in ids])
        data.update({f"jawab_ada_{x}": "0" for x in ids})
        data.update({f"jawab_ada_{hibah}": "3", f"jawab_catatan_{hibah}": "tanpa saksi"})
        c.post(f"/pradaftar/{pid}", data=data)
        from berkas import pradaftar
        k = db.sambung()
        d = pradaftar.muat(k, pid)
        k.close()
        self.assertEqual([x["id"] for x in d["koreksi"]], [hibah])
        self.assertNotIn("Dasar Penguasaan atau Alas Hak",
                         [x["nama"] for x in d["kurang"]])

    def test_terima_bersyarat_menyalin_catatan_koreksi_ke_berkas(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        ids = self._butir(c, pid)
        self._koreksi(c, pid, ids, ids[0], "Belum ditandatangani")
        j = c.post(f"/pradaftar/{pid}/terima", data={"alasan_terima": "menyusul"})
        bid = int(j.headers["Location"].split("/berkas/")[1].split("?")[0])
        k = db.sambung()
        uraian = [r["uraian"] for r in k.execute(
            "SELECT uraian FROM dokumen_pendukung WHERE berkas_id=?", (bid,))]
        k.close()
        self.assertTrue(any("perlu koreksi: Belum ditandatangani" in u for u in uraian),
                        uraian)

    def test_cetakan_memuat_koreksi_beserta_catatannya(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        ids = self._butir(c, pid)
        self._koreksi(c, pid, ids, ids[0], "NIK tidak sesuai")
        formulir = self._docx(c.get(f"/pradaftar/{pid}/docx/kelengkapan"))
        self.assertIn("(perlu koreksi: NIK tidak sesuai)", formulir)
        # yang lengkap kecuali satu koreksi tetap berhak surat pengembalian
        surat = self._docx(c.get(f"/pradaftar/{pid}/docx/pengembalian"))
        self.assertIn("Perlu diperbaiki: NIK tidak sesuai", surat)
        self.assertNotIn("Belum ada", surat)
        self.assertNotIn("{{", surat)

    def test_catatan_koreksi_baku_ditawarkan_dan_bisa_diatur_admin(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        ids = self._butir(c, pid)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn('<datalist id="koreksi-umum">', isi)
        self.assertIn("Belum dilegalisir pejabat yang berwenang", isi)
        # catatan khusus satu butir hanya muncul di datalist butir itu
        j = c.post("/referensi/catatan-koreksi",
                   data={"teks": "SPPT bukan tahun berjalan", "butir_id": str(ids[0]),
                         "aktif": "1"})
        self.assertEqual(j.status_code, 303)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn(f'<datalist id="koreksi-{ids[0]}">', isi)
        self.assertEqual(isi.count("SPPT bukan tahun berjalan"), 1)
        kartu = c.get("/referensi/bagian/catatan_koreksi").get_data(as_text=True)
        self.assertIn("SPPT bukan tahun berjalan", kartu)
        cid = int(re.search(r'name="id" value="(\d+)">\s*<input name="teks" '
                            r'value="SPPT bukan tahun berjalan"', kartu).group(1))
        # dinonaktifkan: tidak ditawarkan lagi
        c.post("/referensi/catatan-koreksi",
               data={"id": str(cid), "teks": "SPPT bukan tahun berjalan",
                     "butir_id": str(ids[0])})
        self.assertNotIn("SPPT bukan tahun berjalan",
                         c.get(f"/pradaftar/{pid}").get_data(as_text=True))
        c.post("/referensi/catatan-koreksi/hapus", data={"id": str(cid)})
        self.assertNotIn("SPPT bukan tahun berjalan",
                         c.get("/referensi/bagian/catatan_koreksi").get_data(as_text=True))

    def test_petugas_tidak_boleh_mengubah_catatan_koreksi_baku(self):
        a = self.klien(masuk=True)
        p = self._petugas(a, "loketkor")
        self._ditolak(p.post("/referensi/catatan-koreksi", data={"teks": "coba"}))
        self._ditolak(p.post("/referensi/catatan-koreksi/hapus", data={"id": "1"}))

    def test_pradaftar_tak_ada_404(self):
        self.assertEqual(self.klien(masuk=True).get("/pradaftar/9999").status_code, 404)


class SusunKelengkapan(Dasar):
    """Menyusun formulir kelengkapan dari layar, dan butir berisi banyak surat.

    Kelas sendiri, basis data sendiri: tiap uji di sini mengubah formulir
    184-HM, dan uji Pradaftar menghitung butirnya.
    """
    _butir = Pradaftar._butir
    _buat = Pradaftar._buat
    _centang = Pradaftar._centang
    _ditolak = Pradaftar._ditolak
    _petugas = Pradaftar._petugas
    _formulir_id = Pradaftar._formulir_id
    ISIAN = Pradaftar.ISIAN

    def _susun(self, c, fid):
        return c.get(f"/referensi/kelengkapan/{fid}").get_data(as_text=True)

    def _id_butir(self, isi, bunyi):
        """id butir di halaman susun menurut bunyinya."""
        m = re.search(r'id="s-(\d+)"[^>]*>(?:(?!</li>).)*?class="susun-nama">'
                      + re.escape(bunyi), isi, re.S)
        return int(m.group(1))

    def _jml_teratas(self, fid):
        k = db.sambung()
        try:
            return k.execute("SELECT COUNT(*) FROM ref_kelengkapan_butir WHERE "
                             "kelengkapan_id=? AND induk_id IS NULL AND aktif=1",
                             (fid,)).fetchone()[0]
        finally:
            k.close()

    def _penanda(self, fid, bunyi):
        k = db.sambung()
        try:
            return k.execute("SELECT penanda FROM ref_kelengkapan_butir "
                             "WHERE kelengkapan_id=? AND nama=?", (fid, bunyi)).fetchone()[0]
        finally:
            k.close()

    def test_halaman_susun_terbuka_dan_petugas_hanya_membaca(self):
        a = self.klien(masuk=True)
        fid = self._formulir_id(a)
        isi = self._susun(a, fid)
        self.assertIn("Susunan butir", isi)
        self.assertIn('id="tambah"', isi)
        p = self._petugas(a, "loketsusun")
        isi = self._susun(p, fid)
        self.assertIn("Dasar Penguasaan atau Alas Hak", isi)
        self.assertNotIn('id="tambah"', isi)
        self._ditolak(p.post(f"/referensi/kelengkapan/{fid}/butir", data={"nama": "x"}))

    def test_tambah_butir_muncul_di_pradaftar_dan_bernomor(self):
        c = self.klien(masuk=True)
        fid = self._formulir_id(c)
        j = c.post(f"/referensi/kelengkapan/{fid}/butir",
                   data={"nama": "Peta Analisis Tata Ruang (PATR)", "sifat": "butir",
                         "wajib": "1"})
        self.assertEqual(j.status_code, 303)
        # masuk di akhir tingkat teratas, nomornya menyusul yang terakhir
        self.assertEqual(self._penanda(fid, "Peta Analisis Tata Ruang (PATR)"),
                         str(self._jml_teratas(fid)))
        pid = self._buat(c)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn("Peta Analisis Tata Ruang (PATR)", isi)
        # butir wajib baru ikut menahan penerimaan
        ids = self._butir(c, pid)
        bid = self._id_butir(self._susun(c, fid), "Peta Analisis Tata Ruang")
        self._centang(c, pid, [x for x in ids if x != bid])
        self.assertIn("belum lengkap", c.get(f"/pradaftar/{pid}").get_data(as_text=True))

    def test_geser_naik_menukar_urutan_dan_nomornya(self):
        c = self.klien(masuk=True)
        fid = self._formulir_id(c)
        bid = self._id_butir(self._susun(c, fid), "Peta Bidang Tanah")
        semula = self._penanda(fid, "Peta Bidang Tanah")
        c.post(f"/referensi/kelengkapan/{fid}/butir/{bid}/geser", data={"arah": "naik"})
        self.assertEqual(int(self._penanda(fid, "Peta Bidang Tanah")), int(semula) - 1)
        # anak-anak alas hak ikut pindah bersama induknya dan tetap berhuruf
        self.assertEqual(self._penanda(fid, "Sertipikat"), "a")
        c.post(f"/referensi/kelengkapan/{fid}/butir/{bid}/geser", data={"arah": "turun"})
        self.assertEqual(self._penanda(fid, "Peta Bidang Tanah"), semula)

    def test_seret_lepas_ke_dalam_kelompok(self):
        c = self.klien(masuk=True)
        fid = self._formulir_id(c)
        c.post(f"/referensi/kelengkapan/{fid}/butir",
               data={"nama": "Butir yang diseret", "sifat": "butir"})
        isi = self._susun(c, fid)
        bid = self._id_butir(isi, "Butir yang diseret")
        grup = self._id_butir(isi, "Dasar Penguasaan atau Alas Hak")
        j = c.post(f"/referensi/kelengkapan/{fid}/butir/{bid}/geser",
                   data={"sasaran": str(grup), "posisi": "dalam"},
                   headers={"X-Diam": "1"})
        self.assertEqual(j.status_code, 204)
        self.assertEqual(self._penanda(fid, "Butir yang diseret"), "k")
        # ke dalam anaknya sendiri ditolak
        j = c.post(f"/referensi/kelengkapan/{fid}/butir/{grup}/geser",
                   data={"sasaran": str(bid), "posisi": "dalam"}, headers={"X-Diam": "1"})
        self.assertEqual(j.status_code, 400)

    def test_hapus_butir_yang_sudah_dijawab_hanya_nonaktif(self):
        c = self.klien(masuk=True)
        fid = self._formulir_id(c)
        c.post(f"/referensi/kelengkapan/{fid}/butir",
               data={"nama": "Butir percobaan hapus", "sifat": "butir", "wajib": "1"})
        c.post(f"/referensi/kelengkapan/{fid}/butir",
               data={"nama": "Butir percobaan dipakai", "sifat": "butir", "wajib": "1"})
        isi = self._susun(c, fid)
        kosong = self._id_butir(isi, "Butir percobaan hapus")
        dipakai = self._id_butir(isi, "Butir percobaan dipakai")
        lama = self._buat(c)
        self._centang(c, lama, [dipakai])
        j = c.post(f"/referensi/kelengkapan/{fid}/butir/{kosong}/hapus")
        self.assertIn("butir_hapus", j.headers["Location"])
        j = c.post(f"/referensi/kelengkapan/{fid}/butir/{dipakai}/hapus")
        self.assertIn("butir_nonaktif", j.headers["Location"])
        # pradaftar lama tetap menampilkannya, yang baru tidak
        self.assertIn("Butir percobaan dipakai",
                      c.get(f"/pradaftar/{lama}").get_data(as_text=True))
        baru = self._buat(c)
        self.assertNotIn("Butir percobaan dipakai",
                         c.get(f"/pradaftar/{baru}").get_data(as_text=True))
        c.post(f"/referensi/kelengkapan/{fid}/butir/{dipakai}/pulihkan")
        self.assertIn("Butir percobaan dipakai",
                      c.get(f"/pradaftar/{baru}").get_data(as_text=True))

    def test_formulir_baru_salinan_dipakai_jenis_haknya(self):
        c = self.klien(masuk=True)
        asal = self._formulir_id(c)
        j = c.post("/referensi/kelengkapan/baru",
                   data={"kode": "hgb-baru", "judul": "Daftar Kelengkapan HGB",
                         "jenis_hak": "HGB", "kegiatan": "*", "salin_dari": str(asal),
                         "aktif": "1", "nomor_otomatis": "1"})
        self.assertEqual(j.status_code, 303)
        fid = int(re.search(r"/kelengkapan/(\d+)", j.headers["Location"]).group(1))
        self.assertIn("Dasar Penguasaan atau Alas Hak", self._susun(c, fid))
        pid = self._buat(c, jenis_hak="HGB")
        self.assertIn("Daftar Kelengkapan HGB", c.get(f"/pradaftar/{pid}").get_data(as_text=True))
        # kode kembar ditolak
        j = c.post("/referensi/kelengkapan/baru",
                   data={"kode": "HGB-BARU", "judul": "lagi", "jenis_hak": "HGB"})
        self.assertIn("galat=", j.headers["Location"])

    def test_bukti_perolehan_lebih_dari_satu_surat(self):
        c = self.klien(masuk=True)
        pid = self._buat(c)
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn("data-tambah-surat", isi)
        j_id = int(re.search(r'name="jawab_uraian_(\d+)"', isi).group(1))
        ids = self._butir(c, pid)
        self._centang(c, pid, ids, **{f"jawab_uraian_{j_id}": [
            "Surat Keterangan Hibah 12/2019", "", "Surat Keterangan Waris 3/2020"]})
        isi = c.get(f"/pradaftar/{pid}").get_data(as_text=True)
        self.assertIn('value="Surat Keterangan Hibah 12/2019"', isi)
        self.assertIn('value="Surat Keterangan Waris 3/2020"', isi)
        # naik jadi berkas: satu surat satu baris dokumen
        j = c.post(f"/pradaftar/{pid}/terima", data={})
        bid = j.headers["Location"].split("/berkas/")[1].split("?")[0]
        isi = c.get(f"/berkas/{bid}").get_data(as_text=True)
        self.assertIn("Surat Keterangan Hibah 12/2019", isi)
        self.assertIn("Surat Keterangan Waris 3/2020", isi)
