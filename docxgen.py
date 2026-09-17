# -*- coding: utf-8 -*-
"""Perakit dokumen DOCX.

Template adalah berkas Word biasa yang dibuat di Microsoft Word, berisi penanda:

    {{nama}}              -> diganti nilai
    {{*daftar.kolom}}     -> paragraf (atau baris tabel) diulang sekali per baris `daftar`
    {{*daftar._akhir}}    -> ";" untuk baris selain terakhir, "." untuk yang terakhir
    {{*daftar._nomor}}    -> nomor urut baris (1, 2, 3, ...)
    {{*daftar._huruf}}    -> a, b, c, ...
    {{*daftar._romawi}}   -> i, ii, iii, ... (penomoran anak butir a.i, a.ii)
    {{?syarat}}           -> paragraf hanya ditulis bila `syarat` bernilai benar
    {{!syarat}}           -> paragraf hanya ditulis bila `syarat` bernilai salah
    {{@foto}}             -> paragrafnya dibuang; foto lapangan dilampirkan di
                             halaman baru pada akhir dokumen
    {{#daftar}} ... {{/daftar}}
                          -> semua yang terapit kedua penanda diulang sekali per
                             baris `daftar`, jadi satu anggota daftar boleh memakai
                             beberapa paragraf atau beberapa baris tabel sekaligus.
                             Di dalamnya, {{?daftar.kolom}} membuang paragraf/baris
                             itu untuk anggota yang kolomnya kosong (mis. NIP).

Format asli (kop, penomoran, tabel, gaya huruf) tidak disentuh sama sekali:
yang diganti hanya teks di dalam run, bukan strukturnya.
"""
import collections
import copy
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as GambarDocx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

PENANDA = re.compile(r"\{\{\s*([*?!@#/]?)\s*([A-Za-z0-9_.]+)\s*\}\}")
BLOK_AWAL = re.compile(r"\{\{\s*#\s*([A-Za-z0-9_]+)\s*\}\}")
BLOK_AKHIR = re.compile(r"\{\{\s*/\s*([A-Za-z0-9_]+)\s*\}\}")

BIARKAN = object()   # dikembalikan pengganti() bila penanda tidak boleh disentuh


# ----------------------------------------------------------------- run
def _runs_paragraf(p):
    return list(p.runs)


def _teks_paragraf(p):
    return "".join(r.text for r in _runs_paragraf(p))


def _ganti_di_paragraf(p, pengganti):
    """Ganti setiap penanda dengan hasil `pengganti(tipe, nama)`.

    Penggantian dilakukan pada run supaya cetak tebal/miring di sekitarnya tetap utuh.
    Penanda yang terpecah ke beberapa run tetap tertangani.
    """
    runs = _runs_paragraf(p)
    if not runs:
        return
    teks = "".join(r.text for r in runs)
    if "{{" not in teks:
        return

    # peta posisi karakter -> (indeks run, offset dalam run)
    peta, pos = [], 0
    for i, r in enumerate(runs):
        for j in range(len(r.text)):
            peta.append((i, j))
        pos += len(r.text)

    cocok = list(PENANDA.finditer(teks))
    isi_run = [r.text for r in runs]

    for m in reversed(cocok):
        nilai = pengganti(m.group(1), m.group(2))
        if nilai is BIARKAN:
            continue
        if nilai is None:
            nilai = ""
        nilai = str(nilai)
        a, b = m.start(), m.end()
        r_awal, o_awal = peta[a]
        r_akhir, o_akhir = peta[b - 1]
        if r_awal == r_akhir:
            s = isi_run[r_awal]
            isi_run[r_awal] = s[:o_awal] + nilai + s[o_akhir + 1:]
        else:
            isi_run[r_awal] = isi_run[r_awal][:o_awal] + nilai
            for i in range(r_awal + 1, r_akhir):
                isi_run[i] = ""
            isi_run[r_akhir] = isi_run[r_akhir][o_akhir + 1:]

    for r, s in zip(runs, isi_run):
        r.text = s


def _semua_paragraf(wadah):
    """Paragraf di badan dokumen dan di dalam seluruh tabel (rekursif)."""
    for p in wadah.paragraphs:
        yield p
    for t in getattr(wadah, "tables", []):
        for baris in t.rows:
            for sel in baris.cells:
                for p in _semua_paragraf(sel):
                    yield p


# ----------------------------------------------------------------- nilai
def _ambil(konteks, nama):
    """Ambil nilai dari konteks; mendukung 'a.b' untuk dict bersarang."""
    nilai = konteks
    for bagian in nama.split("."):
        if isinstance(nilai, dict):
            nilai = nilai.get(bagian)
        else:
            nilai = getattr(nilai, bagian, None)
        if nilai is None:
            return None
    return nilai


def _benar(v):
    if v is None or v is False:
        return False
    if isinstance(v, str):
        return v.strip() not in ("", "-", "0")
    if isinstance(v, (list, tuple, dict)):
        return len(v) > 0
    return bool(v)


# ----------------------------------------------------------------- render
def render(jalur_template, konteks, jalur_keluar, foto=None, kepala_lampiran=None):
    """Rakit `jalur_template` memakai `konteks`, simpan ke `jalur_keluar`.

    foto            = [(jalur_gambar, keterangan), ...] untuk template ber-{{@foto}}
    kepala_lampiran = [(label, nilai), ...] di bawah judul halaman lampiran foto
    """
    dok = Document(jalur_template)
    _proses_wadah(dok, konteks, foto or [], kepala_lampiran or [])
    for bagian in dok.sections:
        for hf in (bagian.header, bagian.footer):
            for p in _semua_paragraf(hf):
                _ganti_di_paragraf(p, lambda t, n: _sederhana(konteks, t, n))
    os.makedirs(os.path.dirname(jalur_keluar), exist_ok=True)
    dok.save(jalur_keluar)
    return jalur_keluar


def _sederhana(konteks, tipe, nama):
    if tipe:
        return ""
    v = _ambil(konteks, nama)
    return "" if v is None else str(v)


def _proses_wadah(dok, konteks, foto, kepala_lampiran):
    # 0) penanda foto dicabut dulu: langkah 4 akan mengosongkannya
    pakai_foto = _cabut_penanda_foto(dok)
    # 1) blok {{#daftar}}..{{/daftar}} dulu: bisa memuat paragraf maupun baris tabel
    _proses_blok(dok, konteks)
    # 2) baris tabel ber-{{*daftar.kolom}} (sekalian membuang tabel yang jadi kosong)
    _proses_baris_tabel(dok, konteks)
    # 3) paragraf berulang dan bersyarat (mengubah jumlah paragraf)
    _proses_ulang_dan_syarat(dok, konteks)
    # 4) substitusi biasa pada semua paragraf yang tersisa
    for p in _semua_paragraf(dok):
        _ganti_di_paragraf(p, lambda t, n: _sederhana(konteks, t, n))
    # 5) foto lapangan, di halaman sendiri pada akhir dokumen
    if pakai_foto:
        _lampiran_foto(dok, foto, kepala_lampiran)


# ------------------------------------------------------- ulangan blok & baris
def _semua_tabel(wadah):
    """Setiap tabel di dalam `wadah`, termasuk tabel di dalam sel (rekursif)."""
    sudah = []                      # elemennya, bukan id-nya: proxy lxml bisa berganti
    tumpuk = [wadah]
    while tumpuk:
        w = tumpuk.pop()
        for t in getattr(w, "tables", []):
            if any(x is t._tbl for x in sudah):
                continue
            sudah.append(t._tbl)
            yield t
            for b in t.rows:
                for sel in b.cells:
                    tumpuk.append(sel)


def _paragraf_di_elemen(el, wadah):
    """Paragraf di dalam satu elemen XML - paragrafnya sendiri atau isi baris/sel."""
    if el.tag == qn("w:p"):
        yield Paragraph(el, wadah)
    for sub in el.findall(".//" + qn("w:p")):
        yield Paragraph(sub, wadah)


def _teks_elemen(el, wadah):
    return "".join(_teks_paragraf(p) for p in _paragraf_di_elemen(el, wadah))


def _buang(el):
    if el.getparent() is not None:
        el.getparent().remove(el)


def _wadah_blok(dok):
    """Elemen yang anak langsungnya boleh diulang: badan dokumen, tabel, dan sel."""
    yield dok.element.body, dok
    sudah = []
    for t in _semua_tabel(dok):
        yield t._tbl, t
        for b in t.rows:
            for sel in b.cells:
                if not any(x is sel._tc for x in sudah):
                    sudah.append(sel._tc)
                    yield sel._tc, sel


def _nilai_kolom(item, kolom):
    if isinstance(item, dict):
        return item.get(kolom)
    return getattr(item, kolom, None)


def _tampil_item(el, wadah, konteks, nama_daftar, item):
    """Benar bila paragraf/baris ini ikut dicetak untuk anggota daftar `item`.

    `{{?daftar.kolom}}` menguji kolom anggotanya, penanda syarat lain menguji konteks.
    """
    for tipe, nama in PENANDA.findall(_teks_elemen(el, wadah)):
        if tipe not in ("?", "!"):
            continue
        if "." in nama and nama.split(".")[0] == nama_daftar:
            v = _benar(_nilai_kolom(item, nama.split(".", 1)[1]))
        else:
            v = _benar(_ambil(konteks, nama))
        if (tipe == "?" and not v) or (tipe == "!" and v):
            return False
    return True


def _salin_kelompok(isi, sebelum, wadah, konteks, nama_daftar, item, i, total):
    """Gandakan `isi` (daftar elemen) untuk satu anggota daftar, sesudah `sebelum`."""
    f = _pembuat_baris(konteks, nama_daftar, item, i, total)
    for el in isi:
        if not _tampil_item(el, wadah, konteks, nama_daftar, item):
            continue
        salinan = copy.deepcopy(el)
        sebelum.addnext(salinan)
        sebelum = salinan
        for p in _paragraf_di_elemen(salinan, wadah):
            _ganti_di_paragraf(p, f)
    return sebelum


def _proses_blok(dok, konteks):
    """Ulang semua yang terapit {{#daftar}} dan {{/daftar}}."""
    for induk, wadah in list(_wadah_blok(dok)):
        for _ in range(200):                      # blok per wadah; sekaligus jaga-jaga
            anak = list(induk)
            awal, nama = None, None
            for i, el in enumerate(anak):
                teks = _teks_elemen(el, wadah)
                m = BLOK_AWAL.search(teks)
                if not m:
                    continue
                if re.search(r"\{\{\s*/\s*" + m.group(1) + r"\s*\}\}", teks):
                    continue      # blok utuh di dalam anak ini - digarap pada wadah dalamnya
                awal, nama = i, m.group(1)
                break
            if awal is None:
                break
            isi, akhir = [], None
            for j in range(awal + 1, len(anak)):
                if BLOK_AKHIR.search(_teks_elemen(anak[j], wadah)):
                    akhir = j
                    break
                isi.append(anak[j])
            if akhir is None:                     # penutupnya lupa ditulis
                if anak[awal].tag == qn("w:p"):
                    _buang(anak[awal])
                else:                             # jangan sampai satu tabel ikut hilang
                    for pp in _paragraf_di_elemen(anak[awal], wadah):
                        _ganti_di_paragraf(pp, lambda t, n: "" if t in ("#", "/") else BIARKAN)
                continue
            baris = _ambil(konteks, nama) or []
            if not isinstance(baris, (list, tuple)):
                baris = []
            sebelum = anak[akhir]
            for i, item in enumerate(baris):
                sebelum = _salin_kelompok(isi, sebelum, wadah, konteks, nama, item, i, len(baris))
            for el in [anak[awal], anak[akhir]] + isi:
                _buang(el)


def _proses_baris_tabel(dok, konteks):
    """Baris tabel ber-{{*daftar.kolom}} digandakan satu baris per anggota daftar."""
    for t in list(_semua_tabel(dok)):
        for b in list(t.rows):
            tr = b._tr
            ulang = [n for tipe, n in PENANDA.findall(_teks_elemen(tr, t)) if tipe == "*"]
            if not ulang:
                continue
            nama = ulang[0].split(".")[0]
            baris = _ambil(konteks, nama) or []
            if not isinstance(baris, (list, tuple)):
                baris = []
            sebelum = tr
            for i, item in enumerate(baris):
                sebelum = _salin_kelompok([tr], sebelum, t, konteks, nama, item, i, len(baris))
            _buang(tr)
    _buang_tabel_kosong(dok)


def _buang_tabel_kosong(dok):
    """Tabel tanpa satu pun baris ditolak Word, jadi tabelnya sekalian dibuang."""
    for t in list(_semua_tabel(dok)):
        if not t._tbl.findall(qn("w:tr")):
            _buang(t._tbl)


def _proses_ulang_dan_syarat(dok, konteks):
    for p in list(_semua_paragraf(dok)):
        teks = _teks_paragraf(p)
        if "{{" not in teks:
            continue
        penanda = PENANDA.findall(teks)
        tipe_ulang = [n for t, n in penanda if t == "*"]
        tipe_syarat = [(t, n) for t, n in penanda if t in ("?", "!")]

        if tipe_syarat:
            tampil = True
            for t, n in tipe_syarat:
                v = _benar(_ambil(konteks, n))
                tampil = tampil and (v if t == "?" else not v)
            if not tampil:
                _hapus_paragraf(p)
                continue
            _ganti_di_paragraf(p, lambda t, n: "" if t in ("?", "!") else BIARKAN)
            teks = _teks_paragraf(p)
            penanda = PENANDA.findall(teks)
            tipe_ulang = [n for t, n in penanda if t == "*"]

        if not tipe_ulang:
            continue

        nama_daftar = tipe_ulang[0].split(".")[0]
        baris = _ambil(konteks, nama_daftar) or []
        if not isinstance(baris, (list, tuple)):
            baris = []
        if not baris:
            _hapus_paragraf(p)
            continue

        induk = p._p
        sebelum = induk
        for i, item in enumerate(baris):
            salinan = copy.deepcopy(induk)
            sebelum.addnext(salinan)
            sebelum = salinan
            pp = Paragraph(salinan, p._parent)
            _ganti_di_paragraf(pp, _pembuat_baris(konteks, nama_daftar, item, i, len(baris)))
        _hapus_paragraf(p)


def _pembuat_baris(konteks, nama_daftar, item, i, total):
    def f(tipe, nama):
        if tipe == "*":
            bagian = nama.split(".", 1)
            kolom = bagian[1] if len(bagian) > 1 else "uraian"
            if kolom == "_akhir":
                return "." if i == total - 1 else ";"
            if kolom == "_nomor":
                return str(i + 1)
            if kolom == "_huruf":
                return chr(ord("a") + i) if i < 26 else str(i + 1)
            if kolom == "_romawi":
                return _romawi(i + 1)
            v = _nilai_kolom(item, kolom)
            return "" if v is None else str(v)
        if tipe:
            return ""
        return _sederhana(konteks, tipe, nama)
    return f


ROMAWI = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]


def _romawi(n):
    """Nomor romawi kecil: i, ii, iii - dipakai penomoran anak butir (a.i, a.ii)."""
    hasil = ""
    for nilai, lambang in ROMAWI:
        while n >= nilai:
            hasil += lambang
            n -= nilai
    return hasil


def _hapus_paragraf(p):
    el = p._p
    if el.getparent() is not None:
        el.getparent().remove(el)


def _cabut_penanda_foto(dok):
    """Buang paragraf yang isinya hanya {{@foto}}. Benar bila template memakainya."""
    ketemu = False
    for p in list(_semua_paragraf(dok)):
        teks = _teks_paragraf(p)
        if "{{@foto" not in teks:
            continue
        ketemu = True
        if not PENANDA.sub("", teks).strip():
            _hapus_paragraf(p)
    return ketemu


def _huruf_utama(dok):
    """Nama huruf yang paling banyak dipakai isi dokumen, supaya lampiran seragam."""
    hitung = collections.Counter(
        r.font.name for p in dok.paragraphs for r in p.runs if r.font.name and r.text.strip())
    if hitung:
        return hitung.most_common(1)[0][0]
    return dok.styles["Normal"].font.name


def _tulis(p, teks, huruf, ukuran, tebal=False):
    r = p.add_run(teks)
    r.bold = tebal
    r.font.size = Pt(ukuran)
    if huruf:
        r.font.name = huruf
    return r


def _rapat(p, sebelum=0, sesudah=0):
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(sebelum), Pt(sesudah)
    f.left_indent = f.first_line_indent = Cm(0)
    f.line_spacing = 1.0
    # gaya bawaan template bisa membawa penomoran - lampiran tidak bernomor
    ppr = p._p.get_or_add_pPr()
    for el in ppr.findall(qn("w:numPr")):
        ppr.remove(el)


def _garis_tabel(tabel):
    tblpr = tabel._tbl.tblPr
    garis = OxmlElement("w:tblBorders")
    for sisi in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{sisi}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        garis.append(el)
    tblpr.append(garis)


def _ukuran_foto(jalur, lebar_maks, tinggi_maks):
    """(lebar, tinggi) EMU yang muat di kotak tanpa mengubah perbandingan sisi."""
    g = GambarDocx.from_file(jalur)
    lebar, tinggi = g.px_width or 1, g.px_height or 1
    skala = min(lebar_maks / lebar, tinggi_maks / tinggi)
    return int(lebar * skala), int(tinggi * skala)


def _buang_ekor_kosong(dok):
    """Paragraf kosong di ujung naskah bisa tumpah jadi halaman putih sebelum
    lampiran. Yang dibuang hanya paragraf tanpa isi sama sekali (cuma pPr)."""
    badan = dok.element.body
    for el in reversed(list(badan)):
        if el.tag == qn("w:sectPr"):
            continue
        if el.tag != qn("w:p"):
            break
        anak = [c for c in el if c.tag != qn("w:pPr")]
        ppr = el.find(qn("w:pPr"))
        if anak or (ppr is not None and ppr.find(qn("w:sectPr")) is not None):
            break
        badan.remove(el)


def _lampiran_foto(dok, foto, kepala):
    """Halaman "Lampiran Dokumentasi Pemeriksaan Lapang": dua foto per baris."""
    foto = [(j, k) for j, k in foto if j and os.path.exists(j)]
    if not foto:
        return
    bagian = dok.sections[-1]
    lebar_teks = bagian.page_width - bagian.left_margin - bagian.right_margin
    kolom = lebar_teks // 2
    lebar_maks = kolom - Cm(0.6)
    tinggi_maks = int(lebar_maks * 0.78)
    huruf = _huruf_utama(dok)
    _buang_ekor_kosong(dok)

    judul = dok.add_paragraph()
    _rapat(judul, sesudah=2)
    judul.paragraph_format.page_break_before = True
    judul.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tulis(judul, "LAMPIRAN DOKUMENTASI PEMERIKSAAN LAPANG", huruf, 12, tebal=True)
    anak = dok.add_paragraph()
    _rapat(anak, sesudah=10)
    anak.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tulis(anak, "Berita Acara Pemeriksaan Tanah oleh Panitia A", huruf, 11)

    kepala = [(a, b) for a, b in kepala if (b or "").strip()]
    for i, (label, nilai) in enumerate(kepala):
        p = dok.add_paragraph()
        _rapat(p, sesudah=8 if i == len(kepala) - 1 else 1)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.2))
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.6))
        p.paragraph_format.left_indent = Cm(4.6)
        p.paragraph_format.first_line_indent = -Cm(4.6)
        _tulis(p, f"{label}	:	{nilai}", huruf, 11)

    baris = (len(foto) + 1) // 2
    tabel = dok.add_table(rows=baris, cols=2)
    tabel.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabel.autofit = False
    _garis_tabel(tabel)
    for r, tr in enumerate(tabel.rows):
        trpr = tr._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))       # satu foto tak terbelah halaman
        for c, sel in enumerate(tr.cells):
            sel.width = kolom
            n = r * 2 + c
            p = sel.paragraphs[0]
            if n >= len(foto):
                continue
            jalur, ket = foto[n]
            # tanpa keterangan, foto berdiri sendiri - tidak ada baris "Foto n"
            ket = (ket or "").strip()
            _rapat(p, sebelum=4, sesudah=2 if ket else 4)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                lebar, tinggi = _ukuran_foto(jalur, lebar_maks, tinggi_maks)
                p.add_run().add_picture(jalur, width=lebar, height=tinggi)
            except Exception:                                      # noqa: BLE001
                _tulis(p, f"[foto tidak terbaca: {os.path.basename(jalur)}]", huruf, 9)
            if ket:
                cap = sel.add_paragraph()
                _rapat(cap, sesudah=4)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _tulis(cap, ket, huruf, 10)
    # Word menuntut paragraf sesudah tabel terakhir dokumen
    _rapat(dok.add_paragraph())


def daftar_penanda(jalur_template):
    """Semua penanda yang dipakai sebuah template - untuk memeriksa kelengkapan konteks."""
    dok = Document(jalur_template)
    hasil = []
    for p in _semua_paragraf(dok):
        for t, n in PENANDA.findall(_teks_paragraf(p)):
            hasil.append((t, n))
    for bagian in dok.sections:
        for hf in (bagian.header, bagian.footer):
            for p in _semua_paragraf(hf):
                for t, n in PENANDA.findall(_teks_paragraf(p)):
                    hasil.append((t, n))
    return hasil
